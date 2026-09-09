#!/usr/bin/env python3
"""Direct DOCX extraction for current RTM automation flows.

This module fills the practical gap left by ``parser.engine.EnhancedParserEngine`` where
``_parse_from_document`` still returns empty populations. It intentionally performs only
source extraction and lightweight classification. Scoring, bidirectional relationship
mapping and higher-order compliance semantics remain separate downstream concerns.

The output shape is compatible with the existing enhanced parser's ``analysis_data`` input:
``sections``, ``rtm_requirements``, ``otc_elements`` and ``del_deliverables``.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document

RTM_ID = re.compile(r"\b(?:RTM|REQ|CReq)[-_ ]?(\d+(?:\.\d+)*)\b", re.IGNORECASE)
OTC_ID = re.compile(r"\b(?:OTC|TC|TEST)[-_ ]?(\d+(?:\.\d+)*)\b", re.IGNORECASE)
DEL_ID = re.compile(r"\b(?:DEL|D)[-_ ]?(\d+(?:\.\d+)*)\b", re.IGNORECASE)
REQUIREMENT_WORDS = re.compile(r"\b(shall|must|required to|requirement)\b", re.IGNORECASE)
TEST_WORDS = re.compile(r"\b(test case|verification|validation|FAT|SAT|expected result)\b", re.IGNORECASE)
DELIVERABLE_WORDS = re.compile(r"\b(deliverable|document|report|drawing|manual|dossier|procedure)\b", re.IGNORECASE)


def _clean(text: str) -> str:
    return " ".join(text.split())


def _stable_id(prefix: str, match: re.Match[str] | None, ordinal: int) -> str:
    if match:
        return f"{prefix}-{match.group(1)}"
    return f"{prefix}-AUTO-{ordinal:04d}"


def _iter_blocks(document: Document):
    """Yield text blocks with source metadata from paragraphs and table rows."""
    for idx, para in enumerate(document.paragraphs, 1):
        text = _clean(para.text)
        if text:
            yield {
                "text": text,
                "source_kind": "paragraph",
                "source_index": idx,
                "style": para.style.name if para.style else "",
            }
    for table_idx, table in enumerate(document.tables, 1):
        for row_idx, row in enumerate(table.rows, 1):
            cells = [_clean(cell.text) for cell in row.cells]
            text = " | ".join(cell for cell in cells if cell)
            if text:
                yield {
                    "text": text,
                    "source_kind": "table_row",
                    "source_index": f"{table_idx}.{row_idx}",
                    "style": "table",
                }


def extract_document(doc_path: str | Path) -> dict[str, Any]:
    """Extract sections and candidate RTM/OTC/DEL populations from a DOCX file."""
    path = Path(doc_path)
    if path.suffix.lower() != ".docx":
        raise ValueError(f"Direct extractor supports .docx only, got: {path.suffix}")
    document = Document(path)

    sections: list[dict[str, Any]] = []
    requirements: list[dict[str, Any]] = []
    tests: list[dict[str, Any]] = []
    deliverables: list[dict[str, Any]] = []
    current_section = path.stem

    req_seen: set[tuple[str, str]] = set()
    test_seen: set[tuple[str, str]] = set()
    del_seen: set[tuple[str, str]] = set()

    for ordinal, block in enumerate(_iter_blocks(document), 1):
        text = block["text"]
        style = block["style"]
        if style.startswith("Heading"):
            current_section = text
            sections.append({
                "id": f"SEC-{len(sections)+1:04d}",
                "title": text,
                "source_kind": block["source_kind"],
                "source_index": block["source_index"],
            })
            continue

        rtm_match = RTM_ID.search(text)
        otc_match = OTC_ID.search(text)
        del_match = DEL_ID.search(text)

        if rtm_match or REQUIREMENT_WORDS.search(text):
            req_id = _stable_id("RTM", rtm_match, ordinal)
            key = (req_id, text)
            if key not in req_seen:
                req_seen.add(key)
                requirements.append({
                    "id": req_id,
                    "description": text,
                    "source_section": current_section,
                    "source_kind": block["source_kind"],
                    "source_index": block["source_index"],
                    "priority": "Unclassified",
                    "verification_method": "Review",
                    "acceptance_criteria": "TBD",
                })

        if otc_match or TEST_WORDS.search(text):
            test_id = _stable_id("OTC", otc_match, ordinal)
            key = (test_id, text)
            if key not in test_seen:
                test_seen.add(key)
                tests.append({
                    "id": test_id,
                    "name": text[:120],
                    "objective": text,
                    "preconditions": "",
                    "test_steps": [],
                    "expected_results": "",
                    "linked_requirements": sorted({m.group(0) for m in RTM_ID.finditer(text)}),
                    "source_section": current_section,
                    "source_kind": block["source_kind"],
                    "source_index": block["source_index"],
                })

        if del_match or DELIVERABLE_WORDS.search(text):
            del_id = _stable_id("DEL", del_match, ordinal)
            key = (del_id, text)
            if key not in del_seen:
                del_seen.add(key)
                deliverables.append({
                    "id": del_id,
                    "name": text[:120],
                    "description": text,
                    "type": "Unclassified",
                    "due_date": "",
                    "responsible_party": "",
                    "status": "TBD",
                    "dependencies": [],
                    "source_section": current_section,
                    "source_kind": block["source_kind"],
                    "source_index": block["source_index"],
                })

    return {
        "source": str(path),
        "sections": sections,
        "rtm_requirements": requirements,
        "otc_elements": tests,
        "del_deliverables": deliverables,
        "counts": {
            "sections": len(sections),
            "rtm_requirements": len(requirements),
            "otc_elements": len(tests),
            "del_deliverables": len(deliverables),
        },
        "classification_boundary": (
            "Candidate extraction only. Presence or lexical classification does not establish "
            "formal compliance, acceptance, verification closure or current QPS authority."
        ),
    }
