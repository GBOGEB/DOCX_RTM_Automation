#!/usr/bin/env python3
"""
Document Converter - Convert DOCX files for RTM processing
"""

import logging
from pathlib import Path
from typing import Dict, Any, Optional
import json
from datetime import datetime

# Set up logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

def run_document_conversion(document_path=None):
    """
    Main function to run document conversion

    Args:
        document_path (str, optional): Path to specific DOCX file to process
                                     If None, will search for files in input directory

    Returns:
        bool: True if successful, False otherwise
    """
    try:
        print(f"🔄 RTM Document Conversion Starting...")

        if document_path:
            # Process specific file
            input_file = Path(document_path)
            if not input_file.exists():
                print(f"❌ File not found: {document_path}")
                return False

            if not input_file.suffix.lower() == '.docx':
                print(f"❌ Not a DOCX file: {document_path}")
                return False

            print(f"📄 Processing: {input_file.name}")

            # Create output directory
            output_dir = Path("output")
            output_dir.mkdir(exist_ok=True)

            # Process the document
            result = process_docx_file(input_file, output_dir)

            if result:
                print(f"✅ Successfully processed: {input_file.name}")
                return True
            else:
                print(f"⚠️  Processing completed with issues: {input_file.name}")
                return False
        else:
            # Process all files in input directory
            input_dir = Path("input")
            if not input_dir.exists():
                print(f"⚠️  Input directory 'input/' not found")
                print(f"   Creating input directory...")
                input_dir.mkdir(exist_ok=True)
                print(f"   Please place DOCX files in the 'input/' directory")
                return False

            docx_files = list(input_dir.glob("*.docx"))
            if not docx_files:
                print(f"⚠️  No DOCX files found in input directory")
                return False

            print(f"📄 Found {len(docx_files)} DOCX files to process")

            # Create output directory
            output_dir = Path("output")
            output_dir.mkdir(exist_ok=True)

            # Process each file
            processed_count = 0
            for docx_file in docx_files:
                print(f"🔄 Processing: {docx_file.name}")
                if process_docx_file(docx_file, output_dir):
                    processed_count += 1

            print(f"✅ Processed {processed_count}/{len(docx_files)} files successfully")
            return processed_count > 0

    except Exception as e:
        print(f"❌ Error in document conversion: {e}")
        return False

def process_docx_file(docx_path, output_dir):
    """
    Process a single DOCX file

    Args:
        docx_path (Path): Path to DOCX file
        output_dir (Path): Output directory

    Returns:
        bool: True if successful
    """
    try:
        from docx import Document

        # Load the document
        doc = Document(docx_path)

        # Extract content
        paragraphs = []
        tables = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                paragraphs.append({
                    'text': paragraph.text.strip(),
                    'style': paragraph.style.name if paragraph.style else 'Normal'
                })

        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    row_data.append(cell.text.strip())
                table_data.append(row_data)
            tables.append(table_data)

        # Create output filename
        output_filename = output_dir / f"{docx_path.stem}_processed.json"

        # Save results
        import json
        result_data = {
            'source_file': str(docx_path),
            'processed_at': datetime.now().isoformat(),
            'paragraphs': paragraphs,
            'tables': tables,
            'stats': {
                'paragraph_count': len(paragraphs),
                'table_count': len(tables)
            }
        }

        with open(output_filename, 'w', encoding='utf-8') as f:
            json.dump(result_data, f, indent=2, ensure_ascii=False)

        print(f"   📊 Extracted {len(paragraphs)} paragraphs and {len(tables)} tables")
        print(f"   💾 Saved to: {output_filename}")

        return True

    except ImportError:
        print(f"   ❌ python-docx library not found")
        print(f"   Install with: pip install python-docx")
        return False
    except Exception as e:
        print(f"   ❌ Error processing {docx_path}: {e}")
        return False

if __name__ == "__main__":
    # Test the document converter
    print("RTM Document Converter Test")
    print("=" * 40)

    # Check dependencies
    try:
        import docx
        print("✅ python-docx is available")
    except ImportError:
        print("⚠️  python-docx not installed")
        print("For full functionality, install with: pip install python-docx")

    # Test with a sample file if available
    test_file = Path("input/MASTER_1805_1144.docx")
    if test_file.exists():
        print(f"\nTesting with: {test_file}")
        result = run_document_conversion(str(test_file))
        print(f"Conversion result: {'Success' if result else 'Failed'}")
    else:
        print(f"\nTest file not found: {test_file}")
        print("Place a DOCX file in the input/ directory to test")
