import os
import sys
import subprocess
import shutil
import logging
from pathlib import Path
import argparse
from docx import Document
import docx_parser
import rtm_generator

#!/usr/bin/env python3
"""
Setup and Test Script for DOCX RTM Automation

This script:
1. Checks the Python environment
2. Installs required dependencies
3. Creates the project directory structure
4. Runs tests to ensure the pipeline functions correctly
"""


# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger('setup_and_test')

# Define project directories
PROJECT_ROOT = Path(__file__).parent.absolute()
DATA_DIR = PROJECT_ROOT / "data"
OUTPUT_DIR = PROJECT_ROOT / "output"
TEST_DIR = PROJECT_ROOT / "tests"
SAMPLE_DIR = DATA_DIR / "samples"

# Define minimum Python version
MIN_PYTHON_VERSION = (3, 7)

def check_python_version():
    """Check if Python version meets the requirements."""
    current_version = sys.version_info[:2]
    
    if current_version < MIN_PYTHON_VERSION:
        logger.error(f"Python {MIN_PYTHON_VERSION[0]}.{MIN_PYTHON_VERSION[1]} or higher is required. "
                   f"You are using Python {current_version[0]}.{current_version[1]}.")
        sys.exit(1)
    
    logger.info(f"Python version check passed: {sys.version.split()[0]}")

def install_dependencies():
    """Install required dependencies using pip."""
    requirements = [
        "python-docx",
        "pandas",
        "openpyxl",
        "pytest",
        "colorama"
    ]
    
    logger.info("Installing dependencies...")
    
    try:
        subprocess.check_call([
            sys.executable, "-m", "pip", "install"] + requirements
        )
        logger.info("Dependencies installed successfully.")
    except subprocess.CalledProcessError:
        logger.error("Failed to install dependencies.")
        sys.exit(1)

def create_directory_structure():
    """Create necessary project directories if they don't exist."""
    directories = [DATA_DIR, OUTPUT_DIR, TEST_DIR, SAMPLE_DIR]
    
    for directory in directories:
        if not directory.exists():
            logger.info(f"Creating directory: {directory}")
            directory.mkdir(parents=True, exist_ok=True)
        else:
            logger.info(f"Directory already exists: {directory}")

def create_sample_files():
    """Create sample files for testing."""
    try:
        
        sample_doc = Document()
        sample_doc.add_heading('Requirements Document', 0)
        
        sample_doc.add_heading('Requirement 1', level=1)
        sample_doc.add_paragraph('REQ-001: The system shall provide user authentication.')
        
        sample_doc.add_heading('Requirement 2', level=1)
        sample_doc.add_paragraph('REQ-002: The system shall encrypt all stored passwords.')
        
        sample_file_path = SAMPLE_DIR / "sample_requirements.docx"
        sample_doc.save(sample_file_path)
        logger.info(f"Created sample document at {sample_file_path}")
        
        return True
    except ImportError:
        logger.warning("Could not create sample DOCX file. Make sure 'python-docx' is installed.")
        return False
    except Exception as e:
        logger.warning(f"Error creating sample file: {e}")
        return False

def test_pipeline():
    """Test the RTM extraction and generation pipeline."""
    try:
        
        logger.info("Testing RTM pipeline...")
        
        # Test with sample file
        sample_file = SAMPLE_DIR / "sample_requirements.docx"
        if not sample_file.exists():
            logger.warning(f"Sample file not found: {sample_file}")
            return False
        
        # Extract requirements
        requirements = docx_parser.extract_requirements(sample_file)
        if not requirements:
            logger.warning("No requirements extracted.")
            return False
        
        logger.info(f"Successfully extracted {len(requirements)} requirements.")
        
        # Generate RTM
        output_file = OUTPUT_DIR / "test_rtm.xlsx"
        rtm_generator.generate_rtm(requirements, output_file)
        
        if output_file.exists():
            logger.info(f"RTM file generated successfully: {output_file}")
            return True
        else:
            logger.warning(f"RTM file was not generated: {output_file}")
            return False
            
    except ImportError as e:
        logger.error(f"Required module not found: {e}")
        return False
    except Exception as e:
        logger.error(f"Test failed with error: {str(e)}")
        return False

def parse_args():
    """Parse command line arguments."""
    parser = argparse.ArgumentParser(description="Setup and test the DOCX RTM Automation project")
    parser.add_argument("--skip-install", action="store_true", help="Skip dependency installation")
    parser.add_argument("--skip-tests", action="store_true", help="Skip running tests")
    return parser.parse_args()

def main():
    """Main function to setup and test the project."""
    args = parse_args()
    
    logger.info("Starting setup and test for DOCX RTM Automation...")
    
    # Check Python version
    check_python_version()
    
    # Create directory structure
    create_directory_structure()
    
    # Install dependencies
    if not args.skip_install:
        install_dependencies()
    else:
        logger.info("Skipping dependency installation")
    
    # Create sample files
    created_samples = create_sample_files()
    
    # Run tests
    if not args.skip_tests:
        if not created_samples:
            logger.warning("Skipping tests because sample files couldn't be created")
        else:
            test_result = test_pipeline()
            if test_result:
                logger.info("All tests completed successfully!")
            else:
                logger.warning("Some tests failed. See the log above for details.")
    else:
        logger.info("Skipping tests")
    
    logger.info("Setup and test completed.")

if __name__ == "__main__":
    main()