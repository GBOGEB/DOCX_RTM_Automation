#!/usr/bin/env python3
"""
Simple Word to Markdown conversion utility
"""
import os
import sys
from pathlib import Path


def convert_docx_to_markdown(docx_file, output_file=None):
    """Convert a DOCX file to Markdown with basic formatting."""
    try:
        from docx import Document
    except ImportError:
        print("Installing python-docx...")
        import subprocess

        subprocess.run([sys.executable, "-m", "pip", "install", "python-docx"])
        from docx import Document

    print(f"Converting {docx_file} to Markdown...")
    doc = Document(docx_file)

    if output_file is None:
        output_file = Path(docx_file).with_suffix(".md")

    # Ensure output directory exists
    output_path = Path(output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with open(output_path, "w", encoding="utf-8") as out_file:
        # Write the title based on the filename
        out_file.write(f"# {Path(docx_file).stem}\n\n")
        out_file.write(f"*Converted from {Path(docx_file).name}*\n\n")

        # Process the document paragraph by paragraph
        for para in doc.paragraphs:
            # Skip empty paragraphs
            if not para.text.strip():
                continue

            # Handle heading styles
            if para.style.name.startswith("Heading"):
                try:
                    level = int(
                        para.style.name[-1]
                    )  # Extract heading level (e.g. Heading 1 -> 1)
                    out_file.write("#" * level + " " + para.text + "\n\n")
                except ValueError:
                    # If we can't extract a level, default to h2
                    out_file.write("## " + para.text + "\n\n")
            # Handle list styles
            elif para.style.name == "List Bullet" or para.style.name == "ListBullet":
                out_file.write("- " + para.text + "\n")
            elif para.style.name == "List Number" or para.style.name == "ListNumber":
                out_file.write("1. " + para.text + "\n")
            # Default: regular paragraph
            else:
                out_file.write(para.text + "\n\n")

        # Process tables simply
        for table in doc.tables:
            # Create header row from first row
            header_row = []
            for cell in table.rows[0].cells:
                header_row.append(cell.text.strip() or " ")
            out_file.write("| " + " | ".join(header_row) + " |\n")

            # Create separator row
            out_file.write("| " + " | ".join(["---"] * len(header_row)) + " |\n")

            # Create data rows
            for row_idx, row in enumerate(table.rows):
                if row_idx == 0:  # Skip header, already handled
                    continue

                row_cells = []
                for cell in row.cells:
                    row_cells.append(cell.text.strip() or " ")
                out_file.write("| " + " | ".join(row_cells) + " |\n")

            out_file.write("\n")

    print(f"Conversion complete! Output saved to {output_path}")
    return output_path


def main():
    """Main function to handle command line arguments."""
    if len(sys.argv) < 2:
        print("Usage: python word_to_md_converter.py <docx_file> [output_md_file]")
        print(
            "\nExample: python word_to_md_converter.py input/document.docx output/document.md"
        )

        # Check if there's a default document to process
        default_docx = Path("input") / "MASTER_1805_1144.docx"
        if default_docx.exists():
            print(f"\nFound default document: {default_docx}")
            docx_file = str(default_docx)
        else:
            # Look for any DOCX file in input directory
            input_dir = Path("input")
            if input_dir.exists():
                docx_files = list(input_dir.glob("*.docx"))
                if docx_files:
                    docx_file = str(docx_files[0])
                    print(f"\nFound DOCX file: {docx_file}")
                else:
                    return 1
            else:
                return 1
    else:
        docx_file = sys.argv[1]

    output_file = sys.argv[2] if len(sys.argv) > 2 else None

    if not os.path.exists(docx_file):
        print(f"Error: File not found - {docx_file}")
        return 1

    # If no output file is specified, create one in the output directory
    if output_file is None:
        output_dir = Path("output")
        output_dir.mkdir(exist_ok=True)
        output_file = output_dir / Path(docx_file).with_suffix(".md").name

    try:
        output_path = convert_docx_to_markdown(docx_file, output_file)
        print(f"\nSuccessfully converted {docx_file} to {output_path}")
        return 0
    except Exception as e:
        print(f"Error during conversion: {e}")
        import traceback

        traceback.print_exc()
        return 1


if __name__ == "__main__":
    sys.exit(main())
