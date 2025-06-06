#!/usr/bin/env python3
"""
Create a test DOCX file with sample requirements for testing the RTM pipeline.
"""

import os
from pathlib import Path


def create_test_document(output_path=None):
    """
    Create a test DOCX file with sample requirements.

    Args:
        output_path: Where to save the file (default: input/test_document.docx)

    Returns:
        Path to the created file
    """
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        print("Installing python-docx...")
        import subprocess
        import sys

        subprocess.run([sys.executable, "-m", "pip", "install", "python-docx"])
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH

    # Create output directory if needed
    if output_path is None:
        input_dir = Path("input")
        input_dir.mkdir(exist_ok=True)
        output_path = input_dir / "test_document.docx"
    else:
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

    # Create a new document
    doc = Document()

    # Add a title
    title = doc.add_heading("Test Document with Requirements", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Add an introduction
    doc.add_paragraph(
        "This is a test document with sample requirements for the RTM pipeline."
    )
    doc.add_paragraph()

    # Add Requirements section
    doc.add_heading("1. Functional Requirements", 1)

    doc.add_heading("FR-1: User Authentication", 2)
    doc.add_paragraph("The system shall provide user authentication capabilities.")

    doc.add_heading("FR-1.1: Login Function", 3)
    doc.add_paragraph("Users must be able to log in with username and password.")

    doc.add_heading("FR-1.2: Password Reset", 3)
    doc.add_paragraph("Users must be able to reset their password.")

    doc.add_heading("FR-2: Data Management", 2)
    doc.add_paragraph("The system shall manage data effectively.")

    # Add Non-Functional Requirements
    doc.add_heading("2. Non-Functional Requirements", 1)

    doc.add_heading("NFR-1: Performance", 2)
    doc.add_paragraph("The system shall respond within 2 seconds.")

    doc.add_heading("NFR-2: Security", 2)
    doc.add_paragraph("User data must be encrypted in transit and at rest.")

    # Add QQQ format requirements to test the new parser
    doc.add_heading("3. Technical Requirements", 1)

    doc.add_heading("QQQ.100: Database Integration", 2)
    doc.add_paragraph("The system shall integrate with SQL and NoSQL databases.")

    doc.add_heading("QQQ.200.10: API Performance", 3)
    doc.add_paragraph("The API shall handle 1000 requests per second.")

    # Add a table with requirements
    doc.add_heading("4. Interface Requirements", 1)
    doc.add_paragraph("The following table defines interface requirements:")

    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"

    # Add table headers
    headers = table.rows[0].cells
    headers[0].text = "Requirement ID"
    headers[1].text = "Description"
    headers[2].text = "Priority"

    # Add table data
    row = table.rows[1].cells
    row[0].text = "IR-1"
    row[1].text = "The system shall provide a REST API"
    row[2].text = "High"

    row = table.rows[2].cells
    row[0].text = "IR-2"
    row[1].text = "The system shall provide a web UI"
    row[2].text = "Medium"

    row = table.rows[3].cells
    row[0].text = "IR-3"
    row[1].text = "The system shall provide a mobile UI"
    row[2].text = "Low"

    # Save the document
    doc.save(output_path)

    print(f"Test document created at {output_path}")
    return output_path


if __name__ == "__main__":
    create_test_document()
