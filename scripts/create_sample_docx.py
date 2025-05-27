from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent


def create_sample_document(filename_str="sample_document.docx"):
    """Creates a sample Word document for RTM automation testing"""

    output_path = BASE_DIR / filename_str
    # Ensure parent directory exists if filename_str includes subdirectories
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # Create a new document
    doc = Document()

    # Add a title
    title = doc.add_heading("Sample Document for RTM Automation", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Add a normal paragraph
    doc.add_paragraph(
        "This is a sample document created for testing the RTM (Requirements Traceability Matrix) Automation tool."
    )

    # Add requirements section
    doc.add_heading("1. Requirements Section", level=1)

    # Add paragraphs with requirement IDs embedded
    p = doc.add_paragraph("Requirement ID: ")
    p.add_run("REQ-001").bold = True
    p.add_run(" - The system shall process input files in DOCX format.")

    p = doc.add_paragraph("Requirement ID: ")
    p.add_run("REQ-002").bold = True
    p.add_run(" - The system shall extract all requirement IDs from the document.")

    p = doc.add_paragraph("Requirement ID: ")
    p.add_run("REQ-003").bold = True
    p.add_run(" - The system shall generate a traceability matrix in Excel format.")

    # Add a table with requirements
    doc.add_heading("2. Requirements Table", level=1)
    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"

    # Add table headers
    header_cells = table.rows[0].cells
    header_cells[0].text = "Requirement ID"
    header_cells[1].text = "Description"
    header_cells[2].text = "Priority"

    # Add table data
    data = [
        ("REQ-004", "The system shall process files within 5 seconds", "High"),
        ("REQ-005", "The system shall support files up to 10MB", "Medium"),
        ("REQ-006", "The system shall log all processing activities", "Low"),
    ]

    for i, (req_id, desc, priority) in enumerate(data):
        row = table.rows[i + 1].cells
        row[0].text = req_id
        row[1].text = desc
        row[2].text = priority

    # Add a page break
    doc.add_page_break()

    # Add content to the new page
    doc.add_heading("3. Additional Requirements", level=1)

    p = doc.add_paragraph("Requirement ID: ")
    p.add_run("REQ-007").bold = True
    p.add_run(" - The system shall provide a user-friendly interface.")

    p = doc.add_paragraph("Requirement ID: ")
    p.add_run("REQ-008").bold = True
    p.add_run(" - The system shall be able to export data in multiple formats.")

    # Save the document
    doc.save(output_path)
    print(f"Sample document created: {output_path}")


if __name__ == "__main__":
    create_sample_document()
