import os
import re
import pandas as pd
import json
from docx import Document
from datetime import datetime
import logging
from pathlib import Path

# Setup basic logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

CONFIG_CACHE = None
BASE_DIR = Path(__file__).resolve().parent

def load_config(config_filename="config.json"):
    """Load configuration settings, caching the result."""
    global CONFIG_CACHE
    if CONFIG_CACHE is not None:
        return CONFIG_CACHE

    config_dir = BASE_DIR / "config"
    config_path = config_dir / config_filename

    default_config = {
        "input_dir": "input",  # Relative to BASE_DIR
        "output_dir": "output", # Relative to BASE_DIR
        "req_doc": "requirements.docx",
        "test_doc": "test_cases.docx",
        "req_pattern": r'(REQ-\d+):\s*(.*)',
        "tc_pattern": r'(TC-\d+):\s*(.*)\s*\[(REQ-\d+)\]',
        "generate_excel": True,
        "generate_markdown": True,
        "generate_html": True
    }

    try:
        if config_path.exists():
            with open(config_path, 'r', encoding='utf-8') as f:
                config = json.load(f)
            # Merge with defaults for any missing keys
            for key in default_config:
                if key not in config:
                    config[key] = default_config[key]
            logger.info(f"Loaded configuration from {config_path}")
        else:
            config = default_config
            # Save default config
            config_dir.mkdir(parents=True, exist_ok=True) # Ensure config directory exists
            with open(config_path, 'w', encoding='utf-8') as f:
                json.dump(default_config, f, indent=4)
            logger.info(f"Created default configuration at {config_path}")
        CONFIG_CACHE = config
        return config
    except Exception as e:
        logger.error(f"Error loading configuration: {e}. Using defaults.", exc_info=True)
        CONFIG_CACHE = default_config
        return default_config

def ensure_directory(directory_path: Path):
    """Create directory if it doesn't exist."""
    if not directory_path.exists():
        directory_path.mkdir(parents=True, exist_ok=True)
        logger.info(f"Created directory: {directory_path}")

def extract_requirements(doc_path: Path):
    """Extract requirements from a Word document."""
    try:
        doc = Document(str(doc_path)) # python-docx might prefer string path
        requirements = []

        config = load_config()
        req_pattern_str = config.get('req_pattern', r'(REQ-\d+):\s*(.*)')
        req_pattern = re.compile(req_pattern_str)

        logger.info(f"--- Processing Requirements from: {doc_path} using pattern: {req_pattern_str} ---")
        for i, paragraph in enumerate(doc.paragraphs):
            paragraph_text = paragraph.text.strip()
            match = req_pattern.search(paragraph_text)
            if match:
                req_id = match.group(1)
                req_desc = match.group(2).strip()
                requirements.append({
                    'ID': req_id,
                    'Description': req_desc
                })

        logger.info(f"Extracted {len(requirements)} requirements from {doc_path}")
        return requirements

    except Exception as e:
        logger.error(f"Error processing requirements document {doc_path}: {e}", exc_info=True)
        return []

def extract_test_cases(doc_path: Path):
    """Extract test cases with requirement references from a Word document."""
    try:
        doc = Document(str(doc_path)) # python-docx might prefer string path
        test_cases = []

        config = load_config()
        tc_pattern_str = config.get('tc_pattern', r'(TC-\d+):\s*(.*)\s*\[(REQ-\d+)\]')
        tc_pattern = re.compile(tc_pattern_str)

        logger.info(f"\n--- Processing Test Cases from: {doc_path} using pattern: {tc_pattern_str} ---")
        for i, paragraph in enumerate(doc.paragraphs):
            paragraph_text = paragraph.text.strip()
            match = tc_pattern.search(paragraph_text)
            if match:
                tc_id = match.group(1)
                tc_desc = match.group(2).strip()
                req_id = match.group(3)
                test_cases.append({
                    'ID': tc_id,
                    'Description': tc_desc,
                    'Requirement': req_id
                })

        logger.info(f"Extracted {len(test_cases)} test cases from {doc_path}")
        return test_cases

    except Exception as e:
        logger.error(f"Error processing test cases document {doc_path}: {e}", exc_info=True)
        return []

def generate_rtm(requirements, test_cases):
    """Generate Requirements Traceability Matrix."""
    rtm_data = []

    for req in requirements:
        req_id = req['ID']
        req_desc = req['Description']

        # Find all test cases that reference this requirement
        linked_tests = [tc for tc in test_cases if tc['Requirement'] == req_id]

        if linked_tests:
            for tc in linked_tests:
                rtm_data.append({
                    'Requirement ID': req_id,
                    'Requirement Description': req_desc,
                    'Test Case ID': tc['ID'],
                    'Test Case Description': tc['Description']
                })
        else:
            # If no test cases reference this requirement
            rtm_data.append({
                'Requirement ID': req_id,
                'Requirement Description': req_desc,
                'Test Case ID': 'N/A',
                'Test Case Description': 'No test coverage'
            })

    return rtm_data

def save_rtm_to_excel(rtm_data, output_path: Path):
    """Save RTM data to Excel file."""
    try:
        ensure_directory(output_path.parent)
        df = pd.DataFrame(rtm_data)
        df.to_excel(output_path, index=False)
        logger.info(f"RTM saved to Excel: {output_path}")
    except Exception as e:
        logger.error(f"Error saving RTM to Excel: {e}", exc_info=True)

def save_rtm_to_markdown(rtm_data, output_path: Path):
    """Save RTM data to Markdown file with improved formatting."""
    try:
        ensure_directory(output_path.parent)
        with open(output_path, 'w', encoding='utf-8') as md_file:
            md_file.write("# Requirements Traceability Matrix\n\n")
            md_file.write(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")

            # Add summary
            req_ids = set([row['Requirement ID'] for row in rtm_data])
            test_ids = set([row['Test Case ID'] for row in rtm_data if row['Test Case ID'] != 'N/A'])
            md_file.write(f"## Summary\n\n")
            md_file.write(f"* **Total Requirements**: {len(req_ids)}\n")
            md_file.write(f"* **Total Test Cases**: {len(test_ids)}\n\n")

            md_file.write("## Traceability Matrix\n\n")
            md_file.write("| Requirement ID | Requirement Description | Test Case ID | Test Case Description |\n")
            md_file.write("|---------------|------------------------|-------------|----------------------|\n")

            for row in rtm_data:
                req_id = row['Requirement ID']
                req_desc = row['Requirement Description'].replace('|', '\\|')  # Escape pipe characters
                tc_id = row['Test Case ID']
                tc_desc = row['Test Case Description'].replace('|', '\\|')  # Escape pipe characters

                # Add hyperlinks for cross-referencing
                req_link = f"[{req_id}](#{req_id.lower()})"
                tc_link = "N/A" if tc_id == 'N/A' else f"[{tc_id}](#{tc_id.lower()})"

                md_file.write(f"| {req_link} | {req_desc} | {tc_link} | {tc_desc} |\n")

            # Add detailed sections
            md_file.write("\n## Requirements Details\n\n")
            for req_id in sorted(req_ids):
                req_data = next((r for r in rtm_data if r['Requirement ID'] == req_id), None)
                if req_data:
                    md_file.write(f"### <a id='{req_id.lower()}'></a>{req_id}\n\n")
                    md_file.write(f"**Description**: {req_data['Requirement Description']}\n\n")

                    # Find all test cases for this requirement
                    related_tests = [r for r in rtm_data if r['Requirement ID'] == req_id and r['Test Case ID'] != 'N/A']
                    if related_tests:
                        md_file.write("**Related Test Cases**:\n\n")
                        for tc in related_tests:
                            md_file.write(f"* [{tc['Test Case ID']}](#{tc['Test Case ID'].lower()}) - {tc['Test Case Description']}\n")
                    else:
                        md_file.write("**Related Test Cases**: None\n")
                    md_file.write("\n")

        logger.info(f"RTM saved to Markdown: {output_path}")
    except Exception as e:
        logger.error(f"Error saving RTM to Markdown: {e}", exc_info=True)

def generate_coverage_report(requirements, test_cases):
    """Generate a test coverage report."""
    total_reqs = len(requirements)
    covered_reqs = len(set([tc['Requirement'] for tc in test_cases]))
    coverage_pct = (covered_reqs / total_reqs) * 100 if total_reqs > 0 else 0

    # Find requirements without coverage
    covered_req_ids = set([tc['Requirement'] for tc in test_cases])
    uncovered_reqs = [req for req in requirements if req['ID'] not in covered_req_ids]

    return {
        'total': total_reqs,
        'covered': covered_reqs,
        'uncovered': total_reqs - covered_reqs,
        'coverage_pct': coverage_pct,
        'uncovered_reqs': uncovered_reqs
    }

def save_coverage_report(coverage_data, output_path: Path):
    """Save coverage report to Markdown file."""
    try:
        ensure_directory(output_path.parent)
        with open(output_path, 'w', encoding='utf-8') as md_file:
            md_file.write("# Test Coverage Report\n\n")
            md_file.write(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")

            md_file.write(f"## Summary\n\n")
            md_file.write(f"* Total Requirements: {coverage_data['total']}\n")
            md_file.write(f"* Covered Requirements: {coverage_data['covered']}\n")
            md_file.write(f"* Uncovered Requirements: {coverage_data['uncovered']}\n")
            md_file.write(f"* Coverage: {coverage_data['coverage_pct']:.2f}%\n\n")

            if coverage_data['uncovered_reqs']:
                md_file.write("## Uncovered Requirements\n\n")
                for req in coverage_data['uncovered_reqs']:
                    md_file.write(f"* **{req['ID']}**: {req['Description']}\n")

        logger.info(f"Coverage report saved to: {output_path}")
    except Exception as e:
        logger.error(f"Error saving coverage report: {e}", exc_info=True)

def generate_html_report(rtm_data, coverage_data, output_path: Path):
    """Generate comprehensive HTML report combining RTM and coverage data."""
    try:
        ensure_directory(output_path.parent)
        with open(output_path, 'w', encoding='utf-8') as html_file:
            html_file.write("""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Requirements Traceability Matrix Report</title>
    <style>
        body { font-family: Arial, sans-serif; margin: 20px; }
        h1, h2, h3 { color: #2c3e50; }
        table { border-collapse: collapse; width: 100%; margin: 20px 0; }
        th, td { border: 1px solid #ddd; padding: 8px; text-align: left; }
        th { background-color: #f2f2f2; }
        tr:nth-child(even) { background-color: #f9f9f9; }
        .summary-box { background-color: #ecf0f1; padding: 15px; border-radius: 5px; margin: 15px 0; }
        .progress-bar-container { width: 100%; background-color: #e0e0e0; border-radius: 4px; }
        .progress-bar { height: 20px; background-color: #27ae60; border-radius: 4px; text-align: center; color: white; }
        .warning { color: #e74c3c; }
    </style>
</head>
<body>
    <h1>Requirements Traceability Matrix Report</h1>
    <p>Generated: """)

            html_file.write(datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
            html_file.write("</p>")

            # Coverage Summary
            html_file.write("""
    <div class="summary-box">
        <h2>Coverage Summary</h2>
        <p><strong>Total Requirements:</strong> """)
            html_file.write(str(coverage_data['total']))
            html_file.write("""</p>
        <p><strong>Covered Requirements:</strong> """)
            html_file.write(str(coverage_data['covered']))
            html_file.write("""</p>
        <p><strong>Coverage Percentage:</strong> """)
            html_file.write(f"{coverage_data['coverage_pct']:.2f}%")
            html_file.write("""</p>
        <div class="progress-bar-container">
            <div class="progress-bar" style="width:""")
            progress_width = min(coverage_data['coverage_pct'], 100)
            html_file.write(f"{progress_width:.2f}%")
            html_file.write(""";">""")
            html_file.write(f"{coverage_data['coverage_pct']:.2f}%")
            html_file.write("""</div>
        </div>
    </div>""")

            # RTM Table
            html_file.write("""
    <h2>Traceability Matrix</h2>
    <table>
        <tr>
            <th>Requirement ID</th>
            <th>Requirement Description</th>
            <th>Test Case ID</th>
            <th>Test Case Description</th>
        </tr>""")

            for row in rtm_data:
                html_file.write("\n        <tr>")
                html_file.write(f"\n            <td>{row['Requirement ID']}</td>")
                html_file.write(f"\n            <td>{row['Requirement Description']}</td>")

                if row['Test Case ID'] == 'N/A':
                    html_file.write(f"\n            <td class='warning'>{row['Test Case ID']}</td>")
                    html_file.write(f"\n            <td class='warning'>{
