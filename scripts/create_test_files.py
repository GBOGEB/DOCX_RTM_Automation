from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.style import WD_STYLE_TYPE

BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"
OUTPUT_DIR = BASE_DIR / "output"

def create_requirements_doc():
    """Create sample requirements document in both MD and DOCX formats"""
    # Create input directory if it doesn't exist
    INPUT_DIR.mkdir(parents=True, exist_ok=True)

    # Sample markdown content
    md_content = """# Requirements Document

## 1. Introduction
This is a sample introduction.

## 2. Functional Requirements

### FR-1: User Authentication
The system shall allow users to authenticate.

#### FR-1.1: Login
The system shall provide a login form.

#### FR-1.2: Logout
The system shall allow users to log out.

### FR-2: Data Management
The system shall manage data effectively.

## 3. Non-Functional Requirements

### NFR-1: Performance
The system shall respond within 2 seconds.
"""

    # Write markdown file
    md_req_path = INPUT_DIR / "requirements.md"
    with open(md_req_path, 'w', encoding='utf-8') as f:
        f.write(md_content)
    print(f"Created sample requirements markdown: {md_req_path}")

    # Create Word document with the same content
    doc = Document()

    # Add heading styles
    styles = doc.styles
    for level in range(1, 5):
        style_name = f'Heading {level}'
        if style_name not in styles:
            style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[style_name]

        font = style.font
        font.size = Pt(16 - level*2)  # Decreasing font size for lower headings

    # Add content to Word document
    doc.add_heading("Requirements Document", 0)

    doc.add_heading("1. Introduction", 1)
    doc.add_paragraph("This is a sample introduction.")

    doc.add_heading("2. Functional Requirements", 1)

    doc.add_heading("FR-1: User Authentication", 2)
    doc.add_paragraph("The system shall allow users to authenticate.")

    doc.add_heading("FR-1.1: Login", 3)
    doc.add_paragraph("The system shall provide a login form.")

    doc.add_heading("FR-1.2: Logout", 3)
    doc.add_paragraph("The system shall allow users to log out.")

    doc.add_heading("FR-2: Data Management", 2)
    doc.add_paragraph("The system shall manage data effectively.")

    doc.add_heading("3. Non-Functional Requirements", 1)

    doc.add_heading("NFR-1: Performance", 2)
    doc.add_paragraph("The system shall respond within 2 seconds.")

    # Save Word document
    docx_req_path = INPUT_DIR / "requirements.docx"
    doc.save(docx_req_path)
    print(f"Created sample requirements document: {docx_req_path}")

def create_test_case_doc():
    """Create sample test case document in DOCX format"""
    # Create input directory if it doesn't exist
    INPUT_DIR.mkdir(parents=True, exist_ok=True)

    # Create Word document
    doc = Document()

    # Add heading styles
    styles = doc.styles
    for level in range(1, 5):
        style_name = f'Heading {level}'
        if style_name not in styles:
            style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[style_name]

        font = style.font
        font.size = Pt(16 - level*2)  # Decreasing font size for lower headings

    # Add content to Word document
    doc.add_heading("Test Cases", 0)

    doc.add_heading("TC-1: Verify User Login", 1)
    doc.add_paragraph("Test case to verify FR-1.1")
    doc.add_paragraph("Steps: Enter username and password, click login")
    doc.add_paragraph("Expected: User is logged in successfully")

    doc.add_heading("TC-2: Verify User Logout", 1)
    doc.add_paragraph("Test case to verify FR-1.2")
    doc.add_paragraph("Steps: Click logout button")
    doc.add_paragraph("Expected: User is logged out successfully")

    doc.add_heading("TC-3: Verify System Performance", 1)
    doc.add_paragraph("Test case to verify NFR-1")
    doc.add_paragraph("Steps: Execute standard operations and measure response time")
    doc.add_paragraph("Expected: All responses occur within 2 seconds")

    # Save Word document
    docx_tc_path = INPUT_DIR / "test_cases.docx"
    doc.save(docx_tc_path)
    print(f"Created sample test cases document: {docx_tc_path}")

def create_test_files():
    """Create additional files and directories needed for testing"""
    # Create output directory if it doesn't exist
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    print(f"Created output directory: {OUTPUT_DIR}")

    # Here you could create any additional files needed for testing
    # Example: placeholder for any configuration files
    sample_config_path = INPUT_DIR / "config.json"
    with open(sample_config_path, 'w', encoding='utf-8') as f:
        f.write('{\n  "version": "1.0"\n}')
    print(f"Created sample config file: {sample_config_path}")

if __name__ == "__main__":
    print("Creating sample files for DOCX RTM Automation...")
    try:
        create_requirements_doc()
        create_test_case_doc()
        create_test_files()
        print("Sample files created successfully. You can now run the RTM automation tool.")
    except ImportError:
        print("Error: Missing required packages. Please install them with:")
        print("pip install python-docx pandas openpyxl")
