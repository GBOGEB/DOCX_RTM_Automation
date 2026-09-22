#!/usr/bin/env python3
"""Build and inspect a visual-style specimen without touching semantic SSOT."""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

REPO = Path(__file__).resolve().parents[1]
BUILDER_PATH = REPO / "scripts" / "build_data_rich_reference_docx.py"


def _load_builder():
    spec = importlib.util.spec_from_file_location("visual_style_builder", BUILDER_PATH)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


builder = _load_builder()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def _set_cell_border(cell, *, left: str | None = None, all_color: str | None = None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    if all_color:
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), all_color)
            borders.append(el)
    if left:
        el = OxmlElement("w:left")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "18")
        el.set(qn("w:color"), left)
        borders.append(el)
    tc_pr.append(borders)


def _font_run(run, name: str, color: str, *, bold: bool = False) -> None:
    run.font.name = name
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.bold = bold


def build_specimen(output: Path, style_path: Path) -> Path:
    visual = builder.load_style(style_path)
    builder.build_reference_docx(output, style_path)
    doc = Document(output)

    doc.add_paragraph("Visible Component Coverage", style="Heading 1")
    doc.add_paragraph(
        "This specimen is presentation-only evidence. It exercises the governed "
        "style roles without altering the semantic source document."
    )

    table = doc.add_table(rows=1, cols=3)
    table.autofit = True
    header = table.rows[0].cells
    for cell, text in zip(header, ("Token", "Role", "Rendered example")):
        cell.text = text
        _set_cell_fill(cell, visual["colors"][visual["tables"]["header_fill_role"]])
        _set_cell_border(cell, all_color=visual["colors"][visual["tables"]["border_role"]])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            _font_run(
                run,
                visual["fonts"]["body"]["name"],
                visual["colors"][visual["tables"]["header_text_role"]],
                bold=True,
            )

    rows = [
        ("NUM", "special number", "1.2.3"),
        ("REQ", "requirement ID", "REQ-042"),
        ("CAP", "caption number", "Figure 3"),
    ]
    for token, role, example in rows:
        cells = table.add_row().cells
        for cell in cells:
            _set_cell_border(cell, all_color=visual["colors"][visual["tables"]["border_role"]])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[0].text = token
        cells[1].text = role
        cells[2].text = example
        for cell in cells:
            for run in cell.paragraphs[0].runs:
                _font_run(run, visual["fonts"]["body"]["name"], visual["colors"]["body"])
        for run in cells[2].paragraphs[0].runs:
            _font_run(
                run,
                visual["fonts"]["body"]["name"],
                visual["colors"]["special_number"],
                bold=True,
            )

    cap = doc.add_paragraph(style=visual["captions"]["paragraph_style"])
    number = cap.add_run("Table 1")
    number.style = visual["captions"]["number_character_style"]
    cap.add_run(" - Governed style token coverage specimen")

    callout = doc.add_table(rows=1, cols=1)
    call_cell = callout.cell(0, 0)
    _set_cell_fill(call_cell, visual["colors"][visual["callouts"]["fill_role"]])
    _set_cell_border(call_cell, left=visual["colors"][visual["callouts"]["accent_role"]])
    p = call_cell.paragraphs[0]
    lead = p.add_run("CONTROL NOTE. ")
    _font_run(
        lead,
        visual["fonts"]["body"]["name"],
        visual["colors"][visual["callouts"]["accent_role"]],
        bold=True,
    )
    body = p.add_run(
        "Visual styling is derived from tokens; semantic authority remains outside this specimen."
    )
    _font_run(
        body,
        visual["fonts"]["body"]["name"],
        visual["colors"][visual["callouts"]["text_role"]],
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def inspect_specimen(docx_path: Path, style_path: Path) -> dict[str, Any]:
    visual = builder.load_style(style_path)
    doc = Document(docx_path)

    paragraph_styles = {p.style.name for p in doc.paragraphs if p.style is not None}
    character_styles = {
        run.style.name
        for p in doc.paragraphs
        for run in p.runs
        if run.style is not None
    }
    required_paragraph = {
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        visual["requirements"]["title_paragraph_style"],
        visual["requirements"]["metadata_paragraph_style"],
        visual["captions"]["paragraph_style"],
    }
    required_character = {
        visual["requirements"]["id_character_style"],
        visual["requirements"]["metadata_label_character_style"],
        visual["captions"]["number_character_style"],
    }

    table_count = len(doc.tables)
    xml = doc._element.xml.upper()
    table_fill = visual["colors"][visual["tables"]["header_fill_role"]].upper()
    callout_fill = visual["colors"][visual["callouts"]["fill_role"]].upper()
    accent = visual["colors"][visual["callouts"]["accent_role"]].upper()

    checks = {
        "paragraph_style_coverage": required_paragraph.issubset(paragraph_styles),
        "character_style_coverage": required_character.issubset(character_styles),
        "table_component_present": table_count >= 2,
        "table_header_fill_present": table_fill in xml,
        "callout_fill_present": callout_fill in xml,
        "callout_accent_present": accent in xml,
        "caption_text_present": any(p.text.startswith("Table 1") for p in doc.paragraphs),
        "control_note_present": any(
            "CONTROL NOTE." in p.text
            for table in doc.tables
            for row in table.rows
            for cell in row.cells
            for p in cell.paragraphs
        ),
    }
    failed = [name for name, ok in checks.items() if not ok]
    return {
        "schema": "docx_rtm.visual_style_specimen_coverage/1.0.0",
        "style_id": visual["style_id"],
        "status": "PASS" if not failed else "FAIL",
        "checks": checks,
        "failed": failed,
        "paragraph_styles_seen": sorted(paragraph_styles),
        "character_styles_seen": sorted(character_styles),
        "table_count": table_count,
        "docx_sha256": sha256_file(docx_path),
        "style_sha256": sha256_file(style_path),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--style-config", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--receipt", type=Path, required=True)
    args = parser.parse_args()

    build_specimen(args.output, args.style_config)
    receipt = inspect_specimen(args.output, args.style_config)
    args.receipt.parent.mkdir(parents=True, exist_ok=True)
    args.receipt.write_text(json.dumps(receipt, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipt, sort_keys=True))
    return 0 if receipt["status"] == "PASS" else 3


if __name__ == "__main__":
    raise SystemExit(main())
