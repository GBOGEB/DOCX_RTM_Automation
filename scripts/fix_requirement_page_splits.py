#!/usr/bin/env python3
"""Layout-aware requirement pagination repair for rendered DOCX files.

The first DOCX/PDF render is treated as measurement evidence. If a governed
requirement block spans PDF pages, this tool inserts a page break before that
requirement title. A second render must then pass --check-only.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
from pathlib import Path
from typing import Any

from docx import Document

REQ_RE = re.compile(r"^\s*(REQ-\d+)\s+[—-]\s+")


class PaginationError(RuntimeError):
    pass


def normalize(text: str) -> str:
    return " ".join(text.split())


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _style_name(paragraph) -> str:
    return paragraph.style.name if paragraph.style is not None else ""


def requirement_blocks(doc: Document) -> list[dict[str, Any]]:
    paragraphs = doc.paragraphs
    blocks: list[dict[str, Any]] = []
    index = 0
    while index < len(paragraphs):
        match = REQ_RE.match(paragraphs[index].text.strip())
        if not match:
            index += 1
            continue
        start = index
        end = index
        cursor = index + 1
        while cursor < len(paragraphs):
            text = paragraphs[cursor].text.strip()
            if _style_name(paragraphs[cursor]) in {"Heading 1", "Heading 2", "Heading 3"}:
                break
            if REQ_RE.match(text):
                break
            if text:
                end = cursor
            cursor += 1
        blocks.append(
            {
                "requirement_id": match.group(1),
                "start": start,
                "end": end,
                "texts": [
                    normalize(paragraphs[i].text)
                    for i in range(start, end + 1)
                    if normalize(paragraphs[i].text)
                ],
            }
        )
        index = max(cursor, index + 1)
    return blocks


def pdf_page_count(pdf_path: Path) -> int:
    result = subprocess.run(
        ["pdfinfo", str(pdf_path)],
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode != 0:
        raise PaginationError(result.stderr.strip() or "pdfinfo failed")
    match = re.search(r"^Pages:\s+(\d+)\s*$", result.stdout, re.MULTILINE)
    if not match:
        raise PaginationError("could not determine PDF page count")
    return int(match.group(1))


def extract_pdf_pages(pdf_path: Path) -> list[str]:
    pages = []
    for page in range(1, pdf_page_count(pdf_path) + 1):
        result = subprocess.run(
            ["pdftotext", "-f", str(page), "-l", str(page), str(pdf_path), "-"],
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode != 0:
            raise PaginationError(result.stderr.strip() or f"pdftotext failed on page {page}")
        pages.append(normalize(result.stdout))
    return pages


def detect_split_requirements(docx_path: Path, page_texts: list[str]) -> list[str]:
    doc = Document(docx_path)
    split_ids: list[str] = []
    for block in requirement_blocks(doc):
        title = block["texts"][0]
        title_pages = [i for i, text in enumerate(page_texts) if title in text]
        if not title_pages:
            rid = block["requirement_id"]
            title_pages = [i for i, text in enumerate(page_texts) if rid in text]
        if not title_pages:
            raise PaginationError(f"{block['requirement_id']}: title not found in rendered PDF")
        page_text = page_texts[title_pages[0]]
        missing = [text for text in block["texts"] if text not in page_text]
        if missing:
            split_ids.append(block["requirement_id"])
    return split_ids


def add_page_breaks(docx_path: Path, requirement_ids: list[str]) -> int:
    wanted = set(requirement_ids)
    if not wanted:
        return 0
    doc = Document(docx_path)
    applied = 0
    for block in requirement_blocks(doc):
        if block["requirement_id"] in wanted:
            doc.paragraphs[block["start"]].paragraph_format.page_break_before = True
            applied += 1
    if applied != len(wanted):
        missing = sorted(wanted - {b["requirement_id"] for b in requirement_blocks(doc)})
        raise PaginationError(f"could not patch requirement(s): {missing}")
    doc.save(docx_path)
    return applied


def evaluate(docx_path: Path, pdf_path: Path, check_only: bool) -> dict[str, Any]:
    page_texts = extract_pdf_pages(pdf_path)
    split_before = detect_split_requirements(docx_path, page_texts)
    applied = 0
    if split_before and not check_only:
        applied = add_page_breaks(docx_path, split_before)
    status = "PASS" if (not split_before or not check_only) else "FAIL"
    return {
        "schema": "docx_rtm.requirement_layout_pagination/1.0.0",
        "mode": "CHECK_ONLY" if check_only else "MEASURE_AND_PATCH",
        "status": status,
        "docx_sha256": sha256_file(docx_path),
        "pdf_sha256": sha256_file(pdf_path),
        "pdf_pages": len(page_texts),
        "split_requirements": split_before,
        "page_breaks_added": applied,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--pdf", type=Path, required=True)
    parser.add_argument("--receipt", type=Path)
    parser.add_argument("--check-only", action="store_true")
    args = parser.parse_args()

    try:
        receipt = evaluate(args.docx, args.pdf, args.check_only)
    except (OSError, PaginationError) as exc:
        print(f"FAIL: {exc}")
        return 2

    if args.receipt:
        args.receipt.parent.mkdir(parents=True, exist_ok=True)
        args.receipt.write_text(json.dumps(receipt, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipt, sort_keys=True))
    if args.check_only and receipt["split_requirements"]:
        return 3
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
