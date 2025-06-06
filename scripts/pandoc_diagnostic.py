#!/usr/bin/env python3
"""
Pandoc Diagnostic Tool

This script checks the Pandoc integration and Lua filter functionality:
1. Verifies that Pandoc is properly installed and configured
2. Tests Lua filters for basic functionality
3. Performs a test conversion to validate the pipeline
"""

import os
import sys
import subprocess
import tempfile
import argparse
import yaml
import json
from pathlib import Path

# Add parent directory to path for imports
parent_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if parent_dir not in sys.path:
    sys.path.insert(0, parent_dir)

# Import from project modules
try:
    from src.modules.pandoc_integration import PandocConverter
except ImportError:
    # Fallback if module not found
    PandocConverter = None
    print("Warning: Could not import PandocConverter from project modules.")


class DiagnosticResult:
    """Container for diagnostic results."""

    def __init__(self):
        self.success = True
        self.tests = []
        self.failures = []
        self.details = {}

    def add_test(self, name, success, message, details=None):
        """Add a test result."""
        self.tests.append(
            {
                "name": name,
                "success": success,
                "message": message,
                "details": details or {},
            }
        )

        if not success:
            self.success = False
            self.failures.append(name)

        if details:
            self.details[name] = details

    def summary(self):
        """Generate a summary of test results."""
        total = len(self.tests)
        passed = sum(1 for test in self.tests if test["success"])

        result = {
            "success": self.success,
            "total": total,
            "passed": passed,
            "failed": total - passed,
            "failures": self.failures,
            "tests": self.tests,
        }

        return result


def check_pandoc_installation():
    """Check if Pandoc is installed and get version information."""
    result = {
        "installed": False,
        "version": None,
        "lua_support": False,
        "filters_support": False,
        "error": None,
    }

    try:
        # Check pandoc version
        proc = subprocess.run(
            ["pandoc", "--version"], capture_output=True, text=True, check=False
        )

        if proc.returncode == 0:
            result["installed"] = True

            # Parse version info
            version_lines = proc.stdout.splitlines()
            if version_lines:
                result["version"] = version_lines[0]

            # Check for Lua support
            result["lua_support"] = "Lua filters" in proc.stdout

            # Check for filter support
            result["filters_support"] = "--lua-filter" in proc.stdout
        else:
            result["error"] = proc.stderr
    except Exception as e:
        result["error"] = str(e)

    return result


def test_lua_filter(filter_path, test_markdown):
    """Test a Lua filter with sample markdown content."""
    result = {"success": False, "output": None, "error": None}

    try:
        # Create temporary files
        with tempfile.NamedTemporaryFile(
            mode="w", suffix=".md", delete=False
        ) as in_file:
            in_file.write(test_markdown)
            input_path = in_file.name

        output_path = input_path + ".output"

        # Run pandoc with the filter
        proc = subprocess.run(
            ["pandoc", input_path, "-o", output_path, "--lua-filter", filter_path],
            capture_output=True,
            text=True,
            check=False,
        )

        if proc.returncode == 0:
            result["success"] = True

            # Read output if it exists
            if os.path.exists(output_path):
                with open(output_path, "r", encoding="utf-8") as f:
                    result["output"] = f.read()

                # Clean up output file
                os.unlink(output_path)
        else:
            result["error"] = proc.stderr

        # Clean up input file
        os.unlink(input_path)
    except Exception as e:
        result["error"] = str(e)

    return result


def test_converter_module():
    """Test the PandocConverter module."""
    result = {
        "available": PandocConverter is not None,
        "import_error": None,
        "initialized": False,
        "error": None,
    }

    if not result["available"]:
        result["import_error"] = "PandocConverter module could not be imported"
        return result

    try:
        # Test initialization
        converter = PandocConverter()
        result["initialized"] = True

        # Test basic method existence
        result["methods"] = {
            "convert_document": hasattr(converter, "convert_document"),
            "extract_rtm_data": hasattr(converter, "extract_rtm_data"),
        }
    except Exception as e:
        result["error"] = str(e)

    return result


def check_lua_files():
    """Check for Lua filter files in the project."""
    config_dir = Path(parent_dir) / "config"
    result = {"found": [], "missing": []}

    # Expected Lua files
    expected_lua = ["extend_headings.lua", "extract_rtm.lua"]

    # Check for expected files
    for lua_file in expected_lua:
        file_path = config_dir / lua_file
        if file_path.exists():
            result["found"].append(str(file_path))
        else:
            result["missing"].append(lua_file)

    # Find any other Lua files
    for lua_file in config_dir.glob("*.lua"):
        if str(lua_file) not in result["found"]:
            result["found"].append(str(lua_file))

    return result


def test_word_to_markdown_conversion(input_file=None):
    """Test Word to Markdown conversion with the project's pipeline."""
    result = {"success": False, "output_file": None, "error": None, "rtm_data": None}

    if not input_file:
        # Create a simple test Word document
        try:
            from docx import Document

            doc = Document()
            doc.add_heading("Test Document", 0)
            doc.add_heading("Section 1", 1)
            doc.add_paragraph("This document contains a test requirement:")
            doc.add_paragraph(
                "The system shall perform validation before saving data. [priority:high]"
            )
            doc.add_heading("Section 2", 1)
            doc.add_paragraph(
                "Another requirement: REQ-001 The system must handle errors gracefully."
            )

            # Save the document to a temporary file
            input_file = tempfile.mktemp(suffix=".docx")
            doc.save(input_file)
            result["input_file"] = input_file
            result["generated_input"] = True
        except ImportError:
            result["error"] = (
                "python-docx is not installed, cannot create test document"
            )
            return result
        except Exception as e:
            result["error"] = f"Failed to create test document: {str(e)}"
            return result
    else:
        result["input_file"] = input_file
        result["generated_input"] = False

    try:
        # Import the conversion function
        from src.core.word_to_md import convert_word_to_md

        # Prepare output path
        output_file = tempfile.mktemp(suffix=".md")
        rtm_output = tempfile.mktemp(suffix=".json")

        # Load configuration
        config_path = os.path.join(parent_dir, "config", "paths.yaml")
        config = {}
        if os.path.exists(config_path):
            with open(config_path, "r", encoding="utf-8") as f:
                config = yaml.safe_load(f)

        # Perform conversion
        success = convert_word_to_md(
            input_file, output_file, config, extract_rtm=True, rtm_output=rtm_output
        )

        result["success"] = success
        result["output_file"] = output_file

        # Check output file
        if os.path.exists(output_file):
            with open(output_file, "r", encoding="utf-8") as f:
                result["output_content"] = f.read()

        # Check RTM data
        if os.path.exists(rtm_output):
            with open(rtm_output, "r", encoding="utf-8") as f:
                result["rtm_data"] = json.load(f)

    except Exception as e:
        result["error"] = str(e)

    # Clean up temporary files
    if result.get("generated_input") and os.path.exists(input_file):
        os.unlink(input_file)

    return result


def run_diagnostics(args):
    """Run full diagnostic suite."""
    results = DiagnosticResult()

    # Check Pandoc installation
    pandoc_result = check_pandoc_installation()
    results.add_test(
        name="pandoc_installation",
        success=pandoc_result["installed"],
        message="Pandoc installation check",
        details=pandoc_result,
    )

    # Check Lua filters
    lua_files_result = check_lua_files()
    lua_files_success = (
        len(lua_files_result["found"]) > 0 and len(lua_files_result["missing"]) == 0
    )
    results.add_test(
        name="lua_filters",
        success=lua_files_success,
        message="Lua filters check",
        details=lua_files_result,
    )

    # Test Lua filter functionality if Pandoc is available
    if (
        pandoc_result["installed"]
        and pandoc_result["lua_support"]
        and lua_files_result["found"]
    ):
        # Test each Lua filter
        for lua_file in lua_files_result["found"]:
            test_markdown = """# Test Document
            
## Section 1

Requirement: REQ-001 Test requirement [priority:high]

## Section 2

The system shall provide authentication.
"""

            lua_test_result = test_lua_filter(lua_file, test_markdown)

            results.add_test(
                name=f"lua_filter_{os.path.basename(lua_file)}",
                success=lua_test_result["success"],
                message=f"Testing Lua filter: {os.path.basename(lua_file)}",
                details=lua_test_result,
            )

    # Test PandocConverter module
    converter_result = test_converter_module()
    results.add_test(
        name="pandoc_converter_module",
        success=converter_result["available"] and converter_result["initialized"],
        message="PandocConverter module check",
        details=converter_result,
    )

    # Test Word to Markdown conversion
    if args.test_conversion:
        conversion_result = test_word_to_markdown_conversion(args.input_file)
        results.add_test(
            name="word_to_md_conversion",
            success=conversion_result["success"],
            message="Word to Markdown conversion test",
            details=conversion_result,
        )

    return results


def main():
    """Main function."""
    parser = argparse.ArgumentParser(
        description="Pandoc and Lua Filter Diagnostic Tool"
    )
    parser.add_argument("--json", action="store_true", help="Output results as JSON")
    parser.add_argument(
        "--test-conversion",
        action="store_true",
        help="Test Word to Markdown conversion",
    )
    parser.add_argument("--input-file", help="Input Word file for conversion test")

    args = parser.parse_args()

    results = run_diagnostics(args)
    summary = results.summary()

    if args.json:
        print(json.dumps(summary, indent=2))
    else:
        # Print human-readable summary
        print("\n===== Pandoc Diagnostic Results =====\n")

        print(f"Total tests: {summary['total']}")
        print(f"Passed: {summary['passed']}")
        print(f"Failed: {summary['failed']}")
        print("")

        # Print detailed test results
        for test in summary["tests"]:
            status = "PASS" if test["success"] else "FAIL"
            # Green or Red
            color = "\033[92m" if test["success"] else "\033[91m"
            print(f"{color}{status}\033[0m: {test['name']} - {test['message']}")

            if not test["success"] and "error" in test["details"]:
                print(f"  Error: {test['details']['error']}")

            print("")

        # Print summary with color
        if summary["success"]:
            print(
                "\033[92mAll tests passed! Pandoc integration is working correctly.\033[0m"
            )
        else:
            print(
                "\033[91mSome tests failed. Please check the detailed results above.\033[0m"
            )

    return 0 if summary["success"] else 1


if __name__ == "__main__":
    sys.exit(main())
