import os
from docx import Document
def create_test_files():
    input_dir = "input"
    os.makedirs(input_dir, exist_ok=True)
    for i in range(3):
        doc = Document()
        doc.add_heading(f"Test Document {i+1}", level=1)
        doc.add_paragraph(f"This is the content of test document {i+1}.")
        doc.save(os.path.join(input_dir, f"test_doc_{i+1}.docx"))

if __name__ == "__main__":
    create_test_files()
    print("Test files created in the 'input' directory.")