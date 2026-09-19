#!/usr/bin/env python3
"""Build the deterministic reference DOCX for data-rich document projections.

Visual choices are loaded from one governed JSON style contract. The style
contract changes appearance only; it does not own semantic content.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any, Dict

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NUM_ID = 77
ABSTRACT_NUM_ID = 77
STYLE_SCHEMA = "docx_rtm.visual_style/1.0.0"
DEFAULT_STYLE_PATH = (
    Path(__file__).resolve().parents[1]
    / "federation"
    / "DATA_RICH_DOCUMENT"
    / "visual_style.json"
)
HEX_RE = re.compile(r"^[0-9A-Fa-f]{6}$")


class StyleConfigError(ValueError):
    pass


def _set_attr(el, name: str, value: str) -> None:
    el.set(qn(name), value)


def _append(parent, tag: str, **attrs):
    el = OxmlElement(tag)
    for key, value in attrs.items():
        _set_attr(el, f"w:{key}", str(value))
    parent.append(el)
    return el


def _remove_children(parent, tag: str) -> None:
    full = qn(tag)
    for child in list(parent):
        if child.tag == full:
            parent.remove(child)


def _require(mapping: Dict[str, Any], key: str) -> Any:
    if key not in mapping:
        raise StyleConfigError(f"visual style missing required key: {key}")
    return mapping[key]


def load_style(path: Path | None = None) -> Dict[str, Any]:
    style_path = path or DEFAULT_STYLE_PATH
    try:
        style = json.loads(style_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise StyleConfigError(f"cannot load visual style {style_path}: {exc}") from exc

    if style.get("schema") != STYLE_SCHEMA:
        raise StyleConfigError(
            f"unsupported visual style schema: {style.get('schema')!r}"
        )
    if not style.get("style_id"):
        raise StyleConfigError("visual style style_id is required")

    for key in ("fonts", "sizes_pt", "colors", "spacing_pt", "page", "numbering"):
        _require(style, key)

    for role, value in style["colors"].items():
        if not isinstance(value, str) or not HEX_RE.match(value):
            raise StyleConfigError(f"invalid hex color for {role}: {value!r}")

    for role, value in style["sizes_pt"].items():
        if not isinstance(value, (int, float)) or value <= 0:
            raise StyleConfigError(f"invalid font size for {role}: {value!r}")
        if abs(float(value) * 2 - round(float(value) * 2)) > 1e-9:
            raise StyleConfigError(
                f"font size for {role} must use Word 0.5 pt granularity: {value!r}"
            )

    if style.get("governance", {}).get("semantic_content_change_allowed") is not False:
        raise StyleConfigError("style governance must forbid semantic content changes")
    if style.get("governance", {}).get("accepted_baseline_mutation_allowed") is not False:
        raise StyleConfigError("style governance must forbid baseline mutation")

    return style


def _font_name(style: Dict[str, Any], role: str) -> str:
    return style["fonts"][role]["name"]


def _color(style: Dict[str, Any], role: str) -> RGBColor:
    return RGBColor.from_string(style["colors"][role])


def _color_hex(style: Dict[str, Any], role: str) -> str:
    return style["colors"][role].upper()


def _get_or_add_style(doc: Document, name: str, style_type: WD_STYLE_TYPE):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, style_type)


def _configure_font(
    target_style,
    *,
    font_name: str,
    size_pt: float,
    color_hex: str,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    target_style.font.name = font_name
    target_style.font.size = Pt(size_pt)
    target_style.font.color.rgb = RGBColor.from_string(color_hex)
    if bold is not None:
        target_style.font.bold = bold
    if italic is not None:
        target_style.font.italic = italic


def _link_style_to_numbering(style, num_id: int, ilvl: int) -> None:
    ppr = style.element.get_or_add_pPr()
    _remove_children(ppr, "w:numPr")
    num_pr = OxmlElement("w:numPr")
    _append(num_pr, "w:ilvl", val=ilvl)
    _append(num_pr, "w:numId", val=num_id)
    ppr.append(num_pr)


def _install_multilevel_numbering(doc: Document, visual: Dict[str, Any]) -> None:
    numbering = doc.part.numbering_part.element

    for child in list(numbering):
        if (
            child.tag == qn("w:abstractNum")
            and child.get(qn("w:abstractNumId")) == str(ABSTRACT_NUM_ID)
        ):
            numbering.remove(child)
        elif child.tag == qn("w:num") and child.get(qn("w:numId")) == str(NUM_ID):
            numbering.remove(child)

    abstract = OxmlElement("w:abstractNum")
    _set_attr(abstract, "w:abstractNumId", str(ABSTRACT_NUM_ID))
    _append(abstract, "w:multiLevelType", val="multilevel")

    level_text = visual["numbering"]["text"]
    number_color = _color_hex(visual, visual["numbering"]["color_role"])
    heading_font = _font_name(visual, "heading")
    heading_sizes = [
        visual["sizes_pt"]["heading_1"],
        visual["sizes_pt"]["heading_2"],
        visual["sizes_pt"]["heading_3"],
    ]

    for ilvl in range(3):
        lvl = OxmlElement("w:lvl")
        _set_attr(lvl, "w:ilvl", str(ilvl))
        _append(lvl, "w:start", val=1)
        _append(lvl, "w:numFmt", val="decimal")
        _append(lvl, "w:pStyle", val=f"Heading{ilvl + 1}")
        _append(lvl, "w:lvlText", val=level_text[ilvl])
        _append(lvl, "w:lvlJc", val="left")
        _append(lvl, "w:suff", val="space")

        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        _set_attr(ind, "w:left", str(ilvl * 360))
        _set_attr(ind, "w:hanging", "0")
        ppr.append(ind)
        lvl.append(ppr)

        rpr = OxmlElement("w:rPr")
        rfonts = OxmlElement("w:rFonts")
        _set_attr(rfonts, "w:ascii", heading_font)
        _set_attr(rfonts, "w:hAnsi", heading_font)
        rpr.append(rfonts)
        _append(rpr, "w:color", val=number_color)
        if visual["numbering"].get("bold") is True:
            _append(rpr, "w:b")
        _append(rpr, "w:sz", val=int(round(float(heading_sizes[ilvl]) * 2)))
        lvl.append(rpr)

        abstract.append(lvl)

    numbering.append(abstract)

    num = OxmlElement("w:num")
    _set_attr(num, "w:numId", str(NUM_ID))
    _append(num, "w:abstractNumId", val=ABSTRACT_NUM_ID)
    numbering.append(num)

    for ilvl, name in enumerate(("Heading 1", "Heading 2", "Heading 3")):
        _link_style_to_numbering(doc.styles[name], NUM_ID, ilvl)


def _configure_styles(doc: Document, visual: Dict[str, Any]) -> None:
    body_font = _font_name(visual, "body")
    heading_font = _font_name(visual, "heading")
    sizes = visual["sizes_pt"]
    colors = visual["colors"]
    spacing = visual["spacing_pt"]

    normal = doc.styles["Normal"]
    _configure_font(
        normal,
        font_name=body_font,
        size_pt=sizes["body"],
        color_hex=colors["body"],
    )
    normal.paragraph_format.space_after = Pt(spacing["body_after"])
    normal.paragraph_format.line_spacing = 1.08

    heading_specs = {
        "Heading 1": (
            sizes["heading_1"],
            colors["heading_primary"],
            spacing["heading_1_before"],
            spacing["heading_1_after"],
        ),
        "Heading 2": (
            sizes["heading_2"],
            colors["heading_secondary"],
            spacing["heading_2_before"],
            spacing["heading_2_after"],
        ),
        "Heading 3": (
            sizes["heading_3"],
            colors["heading_tertiary"],
            spacing["heading_3_before"],
            spacing["heading_3_after"],
        ),
    }
    for name, (size, color, before, after) in heading_specs.items():
        target = doc.styles[name]
        _configure_font(
            target,
            font_name=heading_font,
            size_pt=size,
            color_hex=color,
            bold=True,
        )
        target.paragraph_format.space_before = Pt(before)
        target.paragraph_format.space_after = Pt(after)
        target.paragraph_format.keep_with_next = True

    title = doc.styles["Title"]
    _configure_font(
        title,
        font_name=heading_font,
        size_pt=sizes["title"],
        color_hex=colors["heading_primary"],
        bold=True,
    )
    title.paragraph_format.space_after = Pt(10)

    caption = _get_or_add_style(doc, visual["captions"]["paragraph_style"], WD_STYLE_TYPE.PARAGRAPH)
    caption.base_style = normal
    _configure_font(
        caption,
        font_name=body_font,
        size_pt=sizes["caption"],
        color_hex=colors["caption"],
        italic=bool(visual["captions"].get("italic", True)),
    )
    caption.paragraph_format.space_before = Pt(spacing["caption_before"])
    caption.paragraph_format.space_after = Pt(spacing["caption_after"])
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True

    req_title = _get_or_add_style(
        doc,
        visual["requirements"]["title_paragraph_style"],
        WD_STYLE_TYPE.PARAGRAPH,
    )
    req_title.base_style = normal
    _configure_font(
        req_title,
        font_name=heading_font,
        size_pt=sizes["requirement_title"],
        color_hex=colors["requirement_title"],
        bold=True,
    )
    req_title.paragraph_format.space_before = Pt(spacing["requirement_title_before"])
    req_title.paragraph_format.space_after = Pt(spacing["requirement_title_after"])
    req_title.paragraph_format.keep_with_next = True

    req_meta = _get_or_add_style(
        doc,
        visual["requirements"]["metadata_paragraph_style"],
        WD_STYLE_TYPE.PARAGRAPH,
    )
    req_meta.base_style = normal
    _configure_font(
        req_meta,
        font_name=body_font,
        size_pt=sizes["metadata"],
        color_hex=colors["metadata"],
    )
    req_meta.paragraph_format.space_after = Pt(spacing["metadata_after"])

    req_id = _get_or_add_style(
        doc,
        visual["requirements"]["id_character_style"],
        WD_STYLE_TYPE.CHARACTER,
    )
    _configure_font(
        req_id,
        font_name=heading_font,
        size_pt=sizes["requirement_title"],
        color_hex=colors["requirement_id"],
        bold=True,
    )

    metadata_label = _get_or_add_style(
        doc,
        visual["requirements"]["metadata_label_character_style"],
        WD_STYLE_TYPE.CHARACTER,
    )
    _configure_font(
        metadata_label,
        font_name=body_font,
        size_pt=sizes["metadata"],
        color_hex=colors["special_number"],
        bold=True,
    )

    caption_number = _get_or_add_style(
        doc,
        visual["captions"]["number_character_style"],
        WD_STYLE_TYPE.CHARACTER,
    )
    _configure_font(
        caption_number,
        font_name=body_font,
        size_pt=sizes["caption"],
        color_hex=colors[visual["captions"]["number_color_role"]],
        bold=True,
        italic=bool(visual["captions"].get("italic", True)),
    )


def build_reference_docx(
    output: Path,
    style_path: Path | None = None,
) -> Path:
    visual = load_style(style_path)

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    margins = visual["page"]["margins_in"]
    section.top_margin = Inches(float(margins["top"]))
    section.bottom_margin = Inches(float(margins["bottom"]))
    section.left_margin = Inches(float(margins["left"]))
    section.right_margin = Inches(float(margins["right"]))

    _configure_styles(doc, visual)
    _install_multilevel_numbering(doc, visual)

    # Fixtures make custom styles explicit in the reference package. Pandoc
    # uses the style/numbering parts, not this body, when --reference-doc is
    # applied to the actual governed projection.
    doc.add_paragraph("Data-rich Document Reference", style="Title")
    doc.add_paragraph("Context and Scope", style="Heading 1")
    doc.add_paragraph("Reference template fixture for Heading 1.")
    doc.add_paragraph("Objectives", style="Heading 2")
    doc.add_paragraph("Reference template fixture for Heading 2.")
    doc.add_paragraph("Architecture Decision", style="Heading 1")
    doc.add_paragraph("Section and Content Model", style="Heading 2")
    doc.add_paragraph("Change Transactions", style="Heading 3")
    doc.add_paragraph("Reference template fixture for Heading 3.")

    req = doc.add_paragraph(style=visual["requirements"]["title_paragraph_style"])
    id_run = req.add_run("REQ-001")
    id_run.style = visual["requirements"]["id_character_style"]
    req.add_run(" - Visual style fixture")
    meta = doc.add_paragraph(style=visual["requirements"]["metadata_paragraph_style"])
    label = meta.add_run("Priority:")
    label.style = visual["requirements"]["metadata_label_character_style"]
    meta.add_run(" MUST")

    caption_p = doc.add_paragraph(style=visual["captions"]["paragraph_style"])
    number = caption_p.add_run("Figure 1")
    number.style = visual["captions"]["number_character_style"]
    caption_p.add_run(" - Caption style fixture")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("config/reference.docx"),
    )
    parser.add_argument(
        "--style-config",
        type=Path,
        default=DEFAULT_STYLE_PATH,
    )
    args = parser.parse_args()
    visual = load_style(args.style_config)
    path = build_reference_docx(args.output, args.style_config)
    print(f"PASS: reference DOCX -> {path}")
    print(f"PASS: visual style -> {visual['style_id']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
