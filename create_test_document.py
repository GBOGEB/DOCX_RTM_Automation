#!/usr/bin/env python3
"""
Create Test Document - Generate a sample DOCX file for testing
"""

from pathlib import Path

def create_test_docx():
    """Create a test DOCX file for RTM processing"""
    try:
        from docx import Document

        # Create input directory
        input_dir = Path("input")
        input_dir.mkdir(exist_ok=True)

        # Create a new document
        doc = Document()

        # Add title
        title = doc.add_heading('RTM Test Document', 0)

        # Add some paragraphs
        doc.add_heading('Requirements', level=1)
        doc.add_paragraph('REQ-001: The system shall process DOCX files automatically.')
        doc.add_paragraph('REQ-002: The system shall extract paragraph content.')
        doc.add_paragraph('REQ-003: The system shall extract table data.')

        doc.add_heading('Test Cases', level=1)
        doc.add_paragraph('TC-001: Verify DOCX file loading')
        doc.add_paragraph('TC-002: Verify paragraph extraction')
        doc.add_paragraph('TC-003: Verify table extraction')

        # Add a table
        doc.add_heading('Traceability Matrix', level=1)
        table = doc.add_table(rows=1, cols=3)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Requirement'
        hdr_cells[1].text = 'Test Case'
        hdr_cells[2].text = 'Status'

        # Add table rows
        test_data = [
            ('REQ-001', 'TC-001', 'Pass'),
            ('REQ-002', 'TC-002', 'Pass'),
            ('REQ-003', 'TC-003', 'Pass')
        ]

        for req, tc, status in test_data:
            row_cells = table.add_row().cells
            row_cells[0].text = req
            row_cells[1].text = tc
            row_cells[2].text = status

        # Save the document
        test_file = input_dir / "rtm_test_document.docx"
        doc.save(test_file)

        print(f"✅ Created test document: {test_file}")
        print(f"📊 Document contains:")
        print(f"   • Multiple headings and paragraphs")
        print(f"   • Sample requirements (REQ-001 to REQ-003)")
        print(f"   • Sample test cases (TC-001 to TC-003)")
        print(f"   • Traceability matrix table")

        return True

    except ImportError:
        print(f"❌ python-docx library not found")
        print(f"Install with: pip install python-docx")
        return False
    except Exception as e:
        print(f"❌ Error creating test document: {e}")
        return False

def main():
    """Main function"""
    print("📄 RTM Test Document Creator")
    print("=" * 35)
    print("Creating a sample DOCX file for testing RTM automation...\n")

    if create_test_docx():
        print(f"\n🎯 Next Steps:")
        print(f"1. python main.py                    # Process the test document")
        print(f"2. python find_output_files.py       # Check generated output")
        print(f"3. Review files in output/ directory")
    else:
        print(f"\n❌ Failed to create test document")
        print(f"Please install python-docx: pip install python-docx")

if __name__ == "__main__":
    main()
