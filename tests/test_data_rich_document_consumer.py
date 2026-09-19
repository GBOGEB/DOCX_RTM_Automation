import copy
import hashlib
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document


REPO = Path(__file__).resolve().parents[1]
CONSUMER_PATH = REPO / "src" / "bridges" / "data_rich_document_consumer.py"
BUILDER_PATH = REPO / "scripts" / "build_data_rich_reference_docx.py"
LAYOUT_PATH = REPO / "scripts" / "fix_requirement_page_splits.py"
STYLE_PATH = REPO / "federation" / "DATA_RICH_DOCUMENT" / "visual_style.json"


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


consumer = load_module("data_rich_document_consumer", CONSUMER_PATH)
builder = load_module("build_data_rich_reference_docx", BUILDER_PATH)
layout = load_module("fix_requirement_page_splits", LAYOUT_PATH)


class DataRichDocumentConsumerTests(unittest.TestCase):
    def manifest(self, markdown_hash):
        return {
            "schema": consumer.MANIFEST_SCHEMA,
            "source": {
                "document_id": "ADR-001",
                "document_version": "1.0.0",
                "semantic_ssot": "JSON",
                "git_repo": "GBOGEB/document-organization-system",
                "git_ref": "SOURCE-SHA",
                "path": "contracts/data_rich_document/examples/adr-001.json",
                "sha256": "0" * 64,
            },
            "projection": {
                "format": "MARKDOWN",
                "sha256": markdown_hash,
                "heading_contract": {
                    "maximum_heading_level": 3,
                    "numbering": "TEMPLATE_MANAGED",
                    "heading_text_contains_numbers": False,
                    "style_map": {
                        "1": "Heading 1",
                        "2": "Heading 2",
                        "3": "Heading 3",
                    },
                },
                "section_count": 3,
                "requirement_count": 0,
            },
            "docx_rtm_consumer": {
                "repo": consumer.CONSUMER_REPO,
                "forbidden_numbering_filter": "config/filters/extend_headings.lua",
            },
            "section_map": [
                {
                    "section_id": "SEC-001",
                    "title": "Context",
                    "heading_level": 1,
                },
                {
                    "section_id": "SEC-002",
                    "title": "Details",
                    "heading_level": 2,
                },
                {
                    "section_id": "SEC-003",
                    "title": "Decision",
                    "heading_level": 3,
                },
            ],
            "relations": [],
            "external_references": [],
            "authority": {
                "rendered_output_authoritative": False,
                "authority_transfer": False,
                "engineering_credit": False,
                "compliance_credit": False,
                "semantic_changes_allowed_in_consumer": False,
            },
        }

    def test_reference_builder_has_heading_1_to_3_numbering(self):
        with tempfile.TemporaryDirectory() as tmp:
            ref = Path(tmp) / "reference.docx"
            builder.build_reference_docx(ref, STYLE_PATH)
            proof = consumer.inspect_numbering_contract(ref)
            self.assertEqual(proof["status"], "PASS")
            self.assertEqual(proof["styles"]["Heading1"]["ilvl"], "0")
            self.assertEqual(proof["styles"]["Heading2"]["ilvl"], "1")
            self.assertEqual(proof["styles"]["Heading3"]["ilvl"], "2")
            self.assertEqual(proof["levels"]["0"]["lvl_text"], "%1")
            self.assertEqual(proof["levels"]["1"]["lvl_text"], "%1.%2")
            self.assertEqual(proof["levels"]["2"]["lvl_text"], "%1.%2.%3")


    def test_visual_style_rejects_quarter_point_font_size(self):
        with tempfile.TemporaryDirectory() as tmp:
            style_path = Path(tmp) / "visual_style.json"
            visual = builder.load_style(STYLE_PATH)
            visual["sizes_pt"]["body"] = 10.75
            import json
            style_path.write_text(json.dumps(visual), encoding="utf-8")
            with self.assertRaisesRegex(builder.StyleConfigError, "0.5 pt granularity"):
                builder.load_style(style_path)

    def test_visual_style_contract_is_applied_to_reference_docx(self):
        with tempfile.TemporaryDirectory() as tmp:
            ref = Path(tmp) / "reference.docx"
            visual = builder.load_style(STYLE_PATH)
            builder.build_reference_docx(ref, STYLE_PATH)
            proof = consumer.inspect_visual_style_contract(ref, visual)
            self.assertEqual(proof["status"], "PASS")
            self.assertEqual(proof["style_id"], visual["style_id"])
            self.assertEqual(
                proof["numbering_color"],
                visual["colors"]["special_number"].upper(),
            )
            self.assertEqual(
                proof["styles"]["Caption"]["color"],
                visual["colors"]["caption"].upper(),
            )

    def test_title_rule_is_disabled_and_legacy_blue_border_removed(self):
        with tempfile.TemporaryDirectory() as tmp:
            ref = Path(tmp) / "reference.docx"
            visual = builder.load_style(STYLE_PATH)
            self.assertFalse(visual["title_rule"]["enabled"])
            builder.build_reference_docx(ref, STYLE_PATH)

            styles = consumer._zip_xml(ref, "word/styles.xml")
            title = styles.find(".//w:style[@w:styleId='Title']", consumer.NS)
            self.assertIsNotNone(title)
            self.assertIsNone(title.find("./w:pPr/w:pBdr", consumer.NS))

    def test_visual_semantics_style_requirement_id_metadata_and_caption(self):
        with tempfile.TemporaryDirectory() as tmp:
            ref = Path(tmp) / "reference.docx"
            candidate = Path(tmp) / "candidate.docx"
            visual = builder.load_style(STYLE_PATH)
            builder.build_reference_docx(ref, STYLE_PATH)

            doc = Document(ref)
            doc.add_paragraph("REQ-077 — Styled requirement")
            doc.add_paragraph("Priority: MUST")
            doc.add_paragraph("Figure 7 - Styled caption", style="Caption")
            doc.save(candidate)

            proof = consumer.apply_visual_semantics(candidate, visual)
            self.assertEqual(proof["status"], "PASS")
            self.assertGreaterEqual(proof["requirement_titles"], 1)
            self.assertGreaterEqual(proof["metadata_lines"], 1)
            self.assertGreaterEqual(proof["captions"], 1)

            styled = Document(candidate)
            req = next(p for p in styled.paragraphs if p.text.startswith("REQ-077"))
            self.assertEqual(req.style.name, visual["requirements"]["title_paragraph_style"])
            self.assertEqual(req.runs[0].style.name, visual["requirements"]["id_character_style"])

            meta = next(p for p in styled.paragraphs if p.text.startswith("Priority:"))
            self.assertEqual(meta.style.name, visual["requirements"]["metadata_paragraph_style"])
            self.assertEqual(
                meta.runs[0].style.name,
                visual["requirements"]["metadata_label_character_style"],
            )

            caption = next(p for p in styled.paragraphs if p.text.startswith("Figure 7"))
            self.assertEqual(caption.style.name, visual["captions"]["paragraph_style"])
            self.assertEqual(caption.runs[0].style.name, visual["captions"]["number_character_style"])

    def test_requirement_pagination_applies_structural_hints(self):
        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / "requirement.docx"
            doc = Document()
            doc.add_paragraph("Architecture", style="Heading 1")
            doc.add_paragraph("REQ-002 — Approved iterative refinement")
            doc.add_paragraph(
                "The document model shall record proposed, approved, applied and reverted change sets."
            )
            doc.add_paragraph("Priority: MUST")
            doc.add_paragraph("Risk: HIGH")
            doc.add_paragraph("Rationale: Continuous refinement must remain auditable.")
            doc.add_paragraph("Verification: test / planned")
            doc.add_paragraph("Interfaces", style="Heading 2")
            doc.save(docx_path)

            proof = consumer.enforce_requirement_block_pagination(docx_path)
            self.assertEqual(proof["status"], "PASS")
            self.assertEqual(proof["requirement_block_count"], 1)

            inspection = consumer.inspect_requirement_pagination(docx_path)
            self.assertEqual(inspection["status"], "PASS")
            self.assertEqual(inspection["requirement_block_count"], 1)
            self.assertEqual(inspection["mode"], "KEEP_WITH_NEXT_HINTS")
            self.assertEqual(inspection["observations"][0]["requirement_id"], "REQ-002")

    def test_layout_feedback_detects_split_and_adds_page_break(self):
        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / "layout.docx"
            doc = Document()
            doc.add_paragraph("Architecture", style="Heading 1")
            doc.add_paragraph("REQ-002 — Approved iterative refinement")
            doc.add_paragraph("The document model shall record approved change sets.")
            doc.add_paragraph("Priority: MUST")
            doc.add_paragraph("Risk: HIGH")
            doc.add_paragraph("Interfaces", style="Heading 2")
            doc.save(docx_path)

            page_1 = layout.normalize(
                "Architecture REQ-002 — Approved iterative refinement "
                "The document model shall record approved change sets. Priority: MUST"
            )
            page_2 = layout.normalize("Risk: HIGH Interfaces")
            split = layout.detect_split_requirements(docx_path, [page_1, page_2])
            self.assertEqual(split, ["REQ-002"])

            self.assertEqual(layout.add_page_breaks(docx_path, split), 1)
            repaired = Document(docx_path)
            req = next(p for p in repaired.paragraphs if p.text.startswith("REQ-002"))
            self.assertTrue(req.paragraph_format.page_break_before)

            final_page = layout.normalize(
                "REQ-002 — Approved iterative refinement "
                "The document model shall record approved change sets. "
                "Priority: MUST Risk: HIGH Interfaces"
            )
            self.assertEqual(
                layout.detect_split_requirements(docx_path, ["Architecture", final_page]),
                [],
            )

    def test_manifest_validation_accepts_exact_hash_and_ref(self):
        with tempfile.TemporaryDirectory() as tmp:
            md = Path(tmp) / "input.md"
            md.write_text("# Context\n", encoding="utf-8")
            digest = hashlib.sha256(md.read_bytes()).hexdigest()
            manifest = self.manifest(digest)
            consumer.validate_manifest(manifest, md, "SOURCE-SHA")

    def test_manifest_validation_rejects_markdown_hash_drift(self):
        with tempfile.TemporaryDirectory() as tmp:
            md = Path(tmp) / "input.md"
            md.write_text("# Context\n", encoding="utf-8")
            manifest = self.manifest("f" * 64)
            with self.assertRaisesRegex(consumer.ConsumerError, "hash mismatch"):
                consumer.validate_manifest(manifest, md, "SOURCE-SHA")

    def test_manifest_validation_rejects_source_ref_drift(self):
        with tempfile.TemporaryDirectory() as tmp:
            md = Path(tmp) / "input.md"
            md.write_text("# Context\n", encoding="utf-8")
            digest = hashlib.sha256(md.read_bytes()).hexdigest()
            manifest = self.manifest(digest)
            with self.assertRaisesRegex(consumer.ConsumerError, "source git ref mismatch"):
                consumer.validate_manifest(manifest, md, "OTHER-SHA")

    def test_manifest_validation_rejects_authority_promotion(self):
        with tempfile.TemporaryDirectory() as tmp:
            md = Path(tmp) / "input.md"
            md.write_text("# Context\n", encoding="utf-8")
            digest = hashlib.sha256(md.read_bytes()).hexdigest()
            manifest = self.manifest(digest)
            manifest["authority"]["rendered_output_authoritative"] = True
            with self.assertRaisesRegex(consumer.ConsumerError, "authority guard"):
                consumer.validate_manifest(manifest, md, "SOURCE-SHA")

    def test_manifest_validation_rejects_numbering_filter_ambiguity(self):
        with tempfile.TemporaryDirectory() as tmp:
            md = Path(tmp) / "input.md"
            md.write_text("# Context\n", encoding="utf-8")
            digest = hashlib.sha256(md.read_bytes()).hexdigest()
            manifest = self.manifest(digest)
            manifest["docx_rtm_consumer"]["forbidden_numbering_filter"] = None
            with self.assertRaisesRegex(consumer.ConsumerError, "forbid"):
                consumer.validate_manifest(manifest, md, "SOURCE-SHA")


if __name__ == "__main__":
    unittest.main()
