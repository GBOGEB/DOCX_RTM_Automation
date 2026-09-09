from pathlib import Path

from docx import Document

from parser.direct_docx import extract_analysis_data


def test_direct_docx_extracts_paragraphs_tables_and_context(tmp_path: Path):
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Cryogenic controls", level=1)
    doc.add_paragraph("RTM-317 The Contractor shall provide the required network interface.")
    doc.add_paragraph("OTC-12 Verification test of the network failover sequence.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "DEL-7"
    table.cell(0, 1).text = "Submit the final controls interface report linked to RTM-317."
    doc.save(path)

    data = extract_analysis_data(path)

    assert data["direct_extraction"]["counts"]["sections"] == 1
    assert any(r["id"] == "RTM-317" for r in data["rtm_requirements"])
    assert any(t["id"] == "OTC-12" for t in data["otc_elements"])
    assert any(d["id"] == "DEL-7" for d in data["del_deliverables"])
    row = next(d for d in data["del_deliverables"] if d["id"] == "DEL-7")
    assert row["source_kind"] == "table_row"
    assert "RTM-317" in row["dependencies"]
    assert row["source_section"] == "Cryogenic controls"


def test_direct_docx_rejects_non_docx(tmp_path: Path):
    path = tmp_path / "sample.txt"
    path.write_text("RTM-1 shall exist", encoding="utf-8")

    try:
        extract_analysis_data(path)
    except ValueError as exc:
        assert "only accepts .docx" in str(exc)
    else:
        raise AssertionError("expected ValueError")
