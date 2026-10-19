"""JADE Top21 DOCX->QPS source-contract penetration, synthetic data only.

No engineering authority promotion, no user documents, no external filesystem.
"""
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from parser.direct_docx import extract_analysis_data
from parser.source_extractor import extract_document


def make_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("QPS interface verification", level=1)
    doc.add_paragraph(
        "RTM-317 The Contractor shall provide the helium network interface."
    )
    doc.add_paragraph(
        "OTC-12 Verification of RTM-317 shall demonstrate failover."
    )
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "DEL-7"
    table.cell(0, 1).text = (
        "Submit the interface report linked to RTM-317."
    )
    doc.save(path)


def test_qps_explicit_references_preserved_across_two_extractors(tmp_path):
    path = tmp_path / "synthetic_qps.docx"
    make_docx(path)
    primary = extract_document(path)
    direct = extract_analysis_data(path)
    for family, item_id in (
        ("rtm_requirements", "RTM-317"),
        ("otc_elements", "OTC-12"),
        ("del_deliverables", "DEL-7"),
    ):
        assert any(x["id"] == item_id for x in primary[family])
        assert any(x["id"] == item_id for x in direct[family])
    assert primary["counts"]["sections"] == 1
    assert direct["direct_extraction"]["counts"]["sections"] == 1
    assert any(
        "RTM-317" in x["dependencies"]
        for x in direct["del_deliverables"] if x["id"] == "DEL-7"
    )


def test_clause_only_source_is_not_canonical_rtm(tmp_path):
    path = tmp_path / "unbound_clause.docx"
    doc = Document()
    doc.add_heading("QPS contractual source", level=1)
    doc.add_paragraph(
        "3.3.18.10 The Contractor shall verify cryogenic operating modes."
    )
    doc.save(path)
    source = extract_document(path)
    requirements = source["rtm_requirements"]
    assert len(requirements) == 1
    assert requirements[0]["id"] == "SRC-3.3.18.10"
    assert requirements[0]["canonical_id_state"] == "UNBOUND"
    assert requirements[0]["source_clause"] == "3.3.18.10"


def test_binary_has_valid_ooxml_and_stable_source_digest(tmp_path):
    path = tmp_path / "synthetic_qps.docx"
    make_docx(path)
    original_sha256 = sha256(path.read_bytes()).hexdigest()
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        assert "[Content_Types].xml" in archive.namelist()
        assert "word/document.xml" in archive.namelist()
    first = extract_analysis_data(path)
    second = extract_analysis_data(path)
    assert first == second
    assert sha256(path.read_bytes()).hexdigest() == original_sha256
    assert first["direct_extraction"]["credit_rule"].endswith(
        "not_compliance_credit"
    )


def test_non_docx_suffix_does_not_create_traceability(tmp_path):
    path = tmp_path / "not_a_source.txt"
    path.write_text("RTM-317 shall exist.", encoding="utf-8")
    for parser in (extract_document, extract_analysis_data):
        with pytest.raises(ValueError):
            parser(path)
