from pathlib import Path

from docx import Document

from parser.source_extractor import extract_document


def test_extracts_headings_requirements_tests_deliverables_and_tables(tmp_path: Path):
    doc = Document()
    doc.add_heading("Acceptance", level=1)
    doc.add_paragraph("RTM-516 The Contractor shall verify compressor performance during FAT.")
    doc.add_paragraph("OTC-12 SAT verification shall demonstrate the expected result.")
    doc.add_paragraph("DEL-03 Deliverable report for the acceptance dossier.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "RTM-722"
    table.cell(0, 1).text = "The spare-part warranty must be stated."
    path = tmp_path / "source.docx"
    doc.save(path)

    result = extract_document(path)

    assert result["counts"]["sections"] == 1
    assert any(r["id"] == "RTM-516" for r in result["rtm_requirements"])
    assert any(r["id"] == "RTM-722" for r in result["rtm_requirements"])
    assert any(t["id"] == "OTC-12" for t in result["otc_elements"])
    assert any(d["id"] == "DEL-03" for d in result["del_deliverables"])
    assert any(r["source_kind"] == "table_row" for r in result["rtm_requirements"])


def test_rejects_non_docx(tmp_path: Path):
    path = tmp_path / "source.txt"
    path.write_text("RTM-1 shall do something", encoding="utf-8")
    try:
        extract_document(path)
    except ValueError as exc:
        assert ".docx" in str(exc)
    else:
        raise AssertionError("expected ValueError")
