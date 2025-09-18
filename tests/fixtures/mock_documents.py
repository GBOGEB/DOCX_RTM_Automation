
"""
Mock document generators for testing document processing engines.
"""

import io
import json
from pathlib import Path
from typing import Dict, Any, BinaryIO
from docx import Document
from docx.shared import Inches
import openpyxl
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches as PptxInches
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter

class MockDocumentGenerator:
    """Generate mock documents for testing various processing engines."""
    
    @staticmethod
    def create_word_document(file_path: Path, content_type: str = "basic") -> Path:
        """Create a mock Word document."""
        doc = Document()
        
        if content_type == "basic":
            doc.add_heading('Test Document', 0)
            doc.add_paragraph('This is a test paragraph for document processing.')
            doc.add_heading('Section 1', level=1)
            doc.add_paragraph('Content for section 1 with some formatting.')
            
        elif content_type == "rtm":
            doc.add_heading('Requirements Traceability Matrix', 0)
            doc.add_paragraph('REQ-001: System shall process documents')
            doc.add_paragraph('REQ-002: System shall generate reports')
            doc.add_paragraph('REQ-003: System shall validate inputs')
            
            # Add a table
            table = doc.add_table(rows=1, cols=3)
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Requirement ID'
            hdr_cells[1].text = 'Description'
            hdr_cells[2].text = 'Status'
            
            for i in range(1, 4):
                row_cells = table.add_row().cells
                row_cells[0].text = f'REQ-00{i}'
                row_cells[1].text = f'Requirement {i} description'
                row_cells[2].text = 'Active'
                
        elif content_type == "complex":
            doc.add_heading('Complex Document Structure', 0)
            
            # Multiple sections with different formatting
            for i in range(1, 4):
                doc.add_heading(f'Section {i}', level=1)
                doc.add_paragraph(f'Content for section {i}')
                
                # Add subsections
                for j in range(1, 3):
                    doc.add_heading(f'Subsection {i}.{j}', level=2)
                    doc.add_paragraph(f'Detailed content for subsection {i}.{j}')
                    
            # Add lists
            doc.add_paragraph('Bullet points:')
            for item in ['Item 1', 'Item 2', 'Item 3']:
                doc.add_paragraph(item, style='List Bullet')
        
        doc.save(str(file_path))
        return file_path
    
    @staticmethod
    def create_excel_document(file_path: Path, content_type: str = "basic") -> Path:
        """Create a mock Excel document."""
        wb = openpyxl.Workbook()
        ws = wb.active
        
        if content_type == "basic":
            ws.title = "Test Data"
            ws['A1'] = "Name"
            ws['B1'] = "Value"
            ws['C1'] = "Status"
            
            for i in range(2, 11):
                ws[f'A{i}'] = f"Item {i-1}"
                ws[f'B{i}'] = i * 10
                ws[f'C{i}'] = "Active" if i % 2 == 0 else "Inactive"
                
        elif content_type == "rtm":
            ws.title = "RTM Data"
            headers = ["Requirement ID", "Description", "Priority", "Status", "Test Case"]
            for col, header in enumerate(headers, 1):
                ws.cell(row=1, column=col, value=header)
                ws.cell(row=1, column=col).font = Font(bold=True)
                ws.cell(row=1, column=col).fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            
            for i in range(2, 21):
                ws.cell(row=i, column=1, value=f"REQ-{i-1:03d}")
                ws.cell(row=i, column=2, value=f"Requirement {i-1} description")
                ws.cell(row=i, column=3, value=["High", "Medium", "Low"][(i-2) % 3])
                ws.cell(row=i, column=4, value=["Active", "Draft", "Approved"][(i-2) % 3])
                ws.cell(row=i, column=5, value=f"TC-{i-1:03d}")
                
        elif content_type == "multi_sheet":
            # Create multiple sheets
            for sheet_name in ["Requirements", "Test Cases", "Defects"]:
                if sheet_name != "Sheet":
                    ws = wb.create_sheet(title=sheet_name)
                else:
                    ws.title = sheet_name
                    
                ws['A1'] = f"{sheet_name} Data"
                for i in range(2, 6):
                    ws[f'A{i}'] = f"{sheet_name} Item {i-1}"
                    ws[f'B{i}'] = f"Value {i-1}"
        
        wb.save(str(file_path))
        return file_path
    
    @staticmethod
    def create_powerpoint_document(file_path: Path, content_type: str = "basic") -> Path:
        """Create a mock PowerPoint document."""
        prs = Presentation()
        
        if content_type == "basic":
            # Title slide
            slide_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(slide_layout)
            title = slide.shapes.title
            subtitle = slide.placeholders[1]
            title.text = "Test Presentation"
            subtitle.text = "Document Processing Test"
            
            # Content slide
            slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(slide_layout)
            title = slide.shapes.title
            content = slide.placeholders[1]
            title.text = "Test Content"
            content.text = "This is test content for PowerPoint processing."
            
        elif content_type == "rtm":
            # RTM presentation
            slide_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(slide_layout)
            title = slide.shapes.title
            subtitle = slide.placeholders[1]
            title.text = "Requirements Traceability Matrix"
            subtitle.text = "System Requirements Overview"
            
            # Requirements slide
            slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(slide_layout)
            title = slide.shapes.title
            content = slide.placeholders[1]
            title.text = "System Requirements"
            content.text = "REQ-001: Process documents\nREQ-002: Generate reports\nREQ-003: Validate inputs"
        
        prs.save(str(file_path))
        return file_path
    
    @staticmethod
    def create_pdf_document(file_path: Path, content_type: str = "basic") -> Path:
        """Create a mock PDF document."""
        buffer = io.BytesIO()
        c = canvas.Canvas(buffer, pagesize=letter)
        width, height = letter
        
        if content_type == "basic":
            c.drawString(100, height - 100, "Test PDF Document")
            c.drawString(100, height - 130, "This is a test PDF for document processing.")
            c.drawString(100, height - 160, "Content line 1")
            c.drawString(100, height - 180, "Content line 2")
            c.drawString(100, height - 200, "Content line 3")
            
        elif content_type == "rtm":
            c.drawString(100, height - 100, "Requirements Traceability Matrix")
            c.drawString(100, height - 130, "REQ-001: System shall process documents")
            c.drawString(100, height - 150, "REQ-002: System shall generate reports")
            c.drawString(100, height - 170, "REQ-003: System shall validate inputs")
            
        elif content_type == "multi_page":
            for page in range(3):
                if page > 0:
                    c.showPage()
                c.drawString(100, height - 100, f"Page {page + 1}")
                c.drawString(100, height - 130, f"Content for page {page + 1}")
                for i in range(5):
                    c.drawString(100, height - 160 - (i * 20), f"Line {i + 1} on page {page + 1}")
        
        c.save()
        
        with open(file_path, 'wb') as f:
            f.write(buffer.getvalue())
        
        return file_path
    
    @staticmethod
    def create_markdown_document(file_path: Path, content_type: str = "basic") -> Path:
        """Create a mock Markdown document."""
        content = ""
        
        if content_type == "basic":
            content = """# Test Markdown Document

This is a test markdown document for processing.

## Section 1

Content for section 1 with **bold** and *italic* text.

### Subsection 1.1

- Item 1
- Item 2  
- Item 3

## Section 2

Content for section 2 with a [link](https://example.com).

```python
def test_function():
    return "Hello, World!"
```
"""
        
        elif content_type == "rtm":
            content = """# Requirements Traceability Matrix

## System Requirements

| Requirement ID | Description | Priority | Status |
|----------------|-------------|----------|--------|
| REQ-001 | System shall process documents | High | Active |
| REQ-002 | System shall generate reports | Medium | Active |
| REQ-003 | System shall validate inputs | High | Draft |

## Test Cases

- TC-001: Test document processing
- TC-002: Test report generation
- TC-003: Test input validation
"""
        
        elif content_type == "complex":
            content = """# Complex Document Structure

## Table of Contents

1. [Introduction](#introduction)
2. [Requirements](#requirements)
3. [Implementation](#implementation)
4. [Testing](#testing)

## Introduction

This is a complex markdown document with multiple sections and formatting.

### Purpose

The purpose of this document is to test complex markdown processing.

## Requirements

### Functional Requirements

1. **REQ-001**: System processing capability
2. **REQ-002**: Report generation functionality
3. **REQ-003**: Input validation mechanisms

### Non-Functional Requirements

- Performance: < 5 seconds processing time
- Memory: < 100MB peak usage
- Reliability: 99.9% uptime

## Implementation

```python
class DocumentProcessor:
    def __init__(self):
        self.engines = {}
    
    def process(self, document):
        return self.engines[document.type].process(document)
```

## Testing

Testing includes:

- [ ] Unit tests
- [ ] Integration tests  
- [ ] Performance tests
- [x] Mock data generation
"""
        
        with open(file_path, 'w', encoding='utf-8') as f:
            f.write(content)
        
        return file_path
    
    @staticmethod
    def create_text_document(file_path: Path, content_type: str = "basic") -> Path:
        """Create a mock text document."""
        content = ""
        
        if content_type == "basic":
            content = """Test Text Document

This is a simple text document for testing text processing capabilities.

Section 1
Content for section 1.

Section 2  
Content for section 2 with multiple lines.
Line 2 of section 2.
Line 3 of section 2.
"""
        
        elif content_type == "rtm":
            content = """Requirements Traceability Matrix

REQ-001: System shall process documents
Description: The system must be able to process various document formats
Priority: High
Status: Active

REQ-002: System shall generate reports  
Description: The system must generate comprehensive reports
Priority: Medium
Status: Active

REQ-003: System shall validate inputs
Description: The system must validate all input data
Priority: High
Status: Draft
"""
        
        with open(file_path, 'w', encoding='utf-8') as f:
            f.write(content)
        
        return file_path

def create_all_mock_documents(base_dir: Path) -> Dict[str, Dict[str, Path]]:
    """Create all types of mock documents for testing."""
    base_dir.mkdir(parents=True, exist_ok=True)
    
    documents = {}
    generator = MockDocumentGenerator()
    
    # Word documents
    documents["word"] = {
        "basic": generator.create_word_document(base_dir / "test_basic.docx", "basic"),
        "rtm": generator.create_word_document(base_dir / "test_rtm.docx", "rtm"),
        "complex": generator.create_word_document(base_dir / "test_complex.docx", "complex"),
    }
    
    # Excel documents
    documents["excel"] = {
        "basic": generator.create_excel_document(base_dir / "test_basic.xlsx", "basic"),
        "rtm": generator.create_excel_document(base_dir / "test_rtm.xlsx", "rtm"),
        "multi_sheet": generator.create_excel_document(base_dir / "test_multi.xlsx", "multi_sheet"),
    }
    
    # PowerPoint documents
    documents["powerpoint"] = {
        "basic": generator.create_powerpoint_document(base_dir / "test_basic.pptx", "basic"),
        "rtm": generator.create_powerpoint_document(base_dir / "test_rtm.pptx", "rtm"),
    }
    
    # PDF documents
    documents["pdf"] = {
        "basic": generator.create_pdf_document(base_dir / "test_basic.pdf", "basic"),
        "rtm": generator.create_pdf_document(base_dir / "test_rtm.pdf", "rtm"),
        "multi_page": generator.create_pdf_document(base_dir / "test_multi.pdf", "multi_page"),
    }
    
    # Markdown documents
    documents["markdown"] = {
        "basic": generator.create_markdown_document(base_dir / "test_basic.md", "basic"),
        "rtm": generator.create_markdown_document(base_dir / "test_rtm.md", "rtm"),
        "complex": generator.create_markdown_document(base_dir / "test_complex.md", "complex"),
    }
    
    # Text documents
    documents["text"] = {
        "basic": generator.create_text_document(base_dir / "test_basic.txt", "basic"),
        "rtm": generator.create_text_document(base_dir / "test_rtm.txt", "rtm"),
    }
    
    return documents
