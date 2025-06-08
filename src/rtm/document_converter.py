"""
Document Converter - RTM Automation Module
"""

#!/usr/bin/env python3
"""
Document Converter - Convert DOCX files for RTM processing
"""

import logging
from pathlib import Path
from typing import Dict, Any, Optional
import json
from datetime import datetime

# Set up logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

def run_document_conversion(document_path: str, output_dir: str = "output") -> Dict[str, Any]:
    """
    Convert DOCX document for RTM processing

    Args:
        document_path: Path to the input DOCX file
        output_dir: Directory to save conversion results

    Returns:
        Dictionary with conversion results and metadata
    """
    logger.info(f"Starting document conversion for: {document_path}")

    try:
        # Ensure paths are Path objects
        doc_path = Path(document_path)
        output_path = Path(output_dir)

        # Validate input file
        if not doc_path.exists():
            raise FileNotFoundError(f"Document not found: {document_path}")

        if not doc_path.suffix.lower() == '.docx':
            raise ValueError(f"Expected DOCX file, got: {doc_path.suffix}")

        # Create output directory
        output_path.mkdir(parents=True, exist_ok=True)

        # Prepare conversion results
        conversion_result = {
            "status": "success",
            "input_file": str(doc_path),
            "output_directory": str(output_path),
            "conversion_time": datetime.now().isoformat(),
            "file_size": doc_path.stat().st_size,
            "converted_files": [],
            "metadata": {}
        }

        # Basic document analysis
        logger.info("Analyzing document structure...")
        doc_metadata = analyze_document_structure(doc_path)
        conversion_result["metadata"] = doc_metadata

        # Convert document content
        logger.info("Converting document content...")
        converted_content = convert_document_content(doc_path, output_path)
        conversion_result["converted_files"].extend(converted_content)

        # Save conversion summary
        summary_file = output_path / f"{doc_path.stem}_conversion_summary.json"
        with open(summary_file, 'w', encoding='utf-8') as f:
            json.dump(conversion_result, f, indent=2, default=str)

        conversion_result["converted_files"].append(str(summary_file))

        logger.info(f"Document conversion completed successfully")
        logger.info(f"Generated {len(conversion_result['converted_files'])} output files")

        return conversion_result

    except Exception as e:
        error_result = {
            "status": "error",
            "input_file": document_path,
            "error": str(e),
            "conversion_time": datetime.now().isoformat()
        }

        logger.error(f"Document conversion failed: {e}")
        return error_result

def analyze_document_structure(doc_path: Path) -> Dict[str, Any]:
    """Analyze the structure of a DOCX document"""
    try:
        # Try to import python-docx for proper DOCX handling
        try:
            from docx import Document

            doc = Document(doc_path)

            metadata = {
                "paragraphs_count": len(doc.paragraphs),
                "tables_count": len(doc.tables),
                "has_tables": len(doc.tables) > 0,
                "sections_count": len(doc.sections),
                "analysis_method": "python-docx"
            }

            # Count non-empty paragraphs
            non_empty_paragraphs = sum(1 for p in doc.paragraphs if p.text.strip())
            metadata["non_empty_paragraphs"] = non_empty_paragraphs

            # Analyze tables if present
            if doc.tables:
                table_info = []
                for i, table in enumerate(doc.tables):
                    table_data = {
                        "table_index": i,
                        "rows": len(table.rows),
                        "columns": len(table.columns) if table.rows else 0
                    }
                    table_info.append(table_data)
                metadata["table_details"] = table_info

            logger.info(f"Document analysis: {metadata['paragraphs_count']} paragraphs, {metadata['tables_count']} tables")

        except ImportError:
            # Fallback to basic file analysis
            logger.warning("python-docx not available, using basic analysis")
            metadata = {
                "file_size": doc_path.stat().st_size,
                "file_modified": datetime.fromtimestamp(doc_path.stat().st_mtime).isoformat(),
                "analysis_method": "basic_file_info"
            }

        return metadata

    except Exception as e:
        logger.error(f"Error analyzing document structure: {e}")
        return {
            "error": str(e),
            "analysis_method": "failed"
        }

def convert_document_content(doc_path: Path, output_path: Path) -> list:
    """Convert document content to various formats"""
    converted_files = []

    try:
        # Try to extract content using python-docx
        try:
            from docx import Document

            doc = Document(doc_path)

            # Extract text content
            text_content = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text)

            # Save extracted text
            if text_content:
                text_file = output_path / f"{doc_path.stem}_extracted_text.txt"
                with open(text_file, 'w', encoding='utf-8') as f:
                    f.write('\n'.join(text_content))
                converted_files.append(str(text_file))
                logger.info(f"Extracted text saved to: {text_file}")

            # Extract table data if present
            if doc.tables:
                tables_data = extract_tables_data(doc.tables)
                if tables_data:
                    tables_file = output_path / f"{doc_path.stem}_tables_data.json"
                    with open(tables_file, 'w', encoding='utf-8') as f:
                        json.dump(tables_data, f, indent=2)
                    converted_files.append(str(tables_file))
                    logger.info(f"Table data saved to: {tables_file}")

        except ImportError:
            # Create a placeholder conversion
            logger.warning("python-docx not available, creating placeholder conversion")
            placeholder_file = output_path / f"{doc_path.stem}_conversion_placeholder.txt"
            with open(placeholder_file, 'w', encoding='utf-8') as f:
                f.write(f"Document conversion placeholder for: {doc_path.name}\n")
                f.write(f"Original file size: {doc_path.stat().st_size} bytes\n")
                f.write(f"Conversion time: {datetime.now().isoformat()}\n")
                f.write("\nNote: Install python-docx for full document processing\n")
                f.write("pip install python-docx\n")
            converted_files.append(str(placeholder_file))

        return converted_files

    except Exception as e:
        logger.error(f"Error converting document content: {e}")
        return []

def extract_tables_data(tables) -> list:
    """Extract data from document tables"""
    tables_data = []

    try:
        for table_idx, table in enumerate(tables):
            table_data = {
                "table_index": table_idx,
                "rows": [],
                "row_count": len(table.rows),
                "column_count": len(table.columns) if table.rows else 0
            }

            for row_idx, row in enumerate(table.rows):
                row_data = {
                    "row_index": row_idx,
                    "cells": []
                }

                for cell_idx, cell in enumerate(row.cells):
                    cell_data = {
                        "cell_index": cell_idx,
                        "text": cell.text.strip()
                    }
                    row_data["cells"].append(cell_data)

                table_data["rows"].append(row_data)

            tables_data.append(table_data)

        logger.info(f"Extracted data from {len(tables_data)} tables")
        return tables_data

    except Exception as e:
        logger.error(f"Error extracting table data: {e}")
        return []

def install_dependencies():
    """Check and suggest installation of required dependencies"""
    try:
        import docx
        logger.info("✅ python-docx is available")
        return True
    except ImportError:
        logger.warning("⚠️  python-docx not installed")
        logger.info("For full functionality, install with: pip install python-docx")
        return False

if __name__ == "__main__":
    # Test the document converter
    print("RTM Document Converter Test")
    print("=" * 40)

    # Check dependencies
    install_dependencies()

    # Test with a sample file if available
    test_file = Path("input/MASTER_1805_1144.docx")
    if test_file.exists():
        print(f"\nTesting with: {test_file}")
        result = run_document_conversion(str(test_file))
        print(f"Conversion result: {result['status']}")
        if result['status'] == 'success':
            print(f"Output files: {len(result['converted_files'])}")
        else:
            print(f"Error: {result.get('error', 'Unknown')}")
    else:
        print(f"\nTest file not found: {test_file}")
        print("Place a DOCX file in the input/ directory to test")
