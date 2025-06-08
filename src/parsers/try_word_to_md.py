#!/usr/bin/env python3
"""
Improved Word to Markdown converter.
Handles headers, formatting, lists, and tables.
"""

import os
import sys
import re
from pathlib import Path


def convert_docx_to_markdown(docx_file, output_file=None):
    """Convert a DOCX file to Markdown with better formatting."""
    try:
        from docx import Document
    except ImportError:
        print("Installing python-docx...")
        import subprocess

        subprocess.run([sys.executable, "-m", "pip", "install", "python-docx"])
        from docx import Document

    print(f"Converting {docx_file} to Markdown...")
    doc = Document(docx_file)

    if output_file is None:
        output_file = Path(docx_file).with_suffix(".md")

    # Ensure output directory exists
    os.makedirs(os.path.dirname(os.path.abspath(output_file)), exist_ok=True)

    with open(output_file, "w", encoding="utf-8") as out_file:
        # Process the document paragraph by paragraph
        for para in doc.paragraphs:
            # Skip empty paragraphs
            if not para.text.strip():
                out_file.write("\n")
                continue

            # Handle heading styles
            if para.style.name.startswith("Heading"):
                try:
                    # Extract heading level more safely
                    match = re.search(r"(\d+)$", para.style.name)
                    if match:
                        level = int(match.group(1))
                    else:
                        # Handle custom heading styles that don't end with a number
                        # Determine level based on style name or default to level 2
                        if "title" in para.style.name.lower():
                            level = 1
                        else:
                            level = 2
                        print(
                            f"Warning: Found non-standard heading style: '{para.style.name}', using level {level}"
                        )

                    out_file.write("#" * level + " " + para.text + "\n\n")
                except (ValueError, AttributeError) as e:
                    # Fallback to heading level 2 in case of any error
                    print(
                        f"Warning: Error processing heading style '{para.style.name}': {e}"
                    )
                    out_file.write("## " + para.text + "\n\n")
            # Handle list styles
            elif para.style.name == "List Bullet":
                out_file.write("- " + para.text + "\n")
            elif para.style.name == "List Number":
                out_file.write("1. " + para.text + "\n")
            # Default: regular paragraph
            else:
                # Better handling of paragraph formatting
                text = ""
                for run in para.runs:
                    run_text = run.text
                    if run.bold and run.italic:
                        run_text = f"***{run_text}***"
                    elif run.bold:
                        run_text = f"**{run_text}**"
                    elif run.italic:
                        run_text = f"*{run_text}*"
                    elif run.underline:
                        # Check if it might be a URL
                        if run.text.lower().startswith(("http://", "https://")):
                            run_text = f"[{run_text}]({run_text})"
                        else:
                            run_text = f"__{run_text}__"
                    text += run_text

                out_file.write(text + "\n\n")

        # Process tables with better formatting
        for table in doc.tables:
            # Create header row
            header_row = []
            for cell in table.rows[0].cells:
                header_row.append(cell.text.strip() or " ")
            out_file.write("| " + " | ".join(header_row) + " |\n")

            # Create separator row
            out_file.write("| " + " | ".join(["---"] * len(header_row)) + " |\n")

            # Create data rows
            for row_idx, row in enumerate(table.rows):
                if row_idx == 0:  # Skip header, already handled
                    continue

                row_cells = []
                for cell in row.cells:
                    row_cells.append(cell.text.strip() or " ")
                out_file.write("| " + " | ".join(row_cells) + " |\n")

            out_file.write("\n")

    print(f"Conversion complete! Output saved to {output_file}")
    return output_file


def main():
    """Main function to handle command line arguments."""
    if len(sys.argv) < 2:
        # Look for test_document.docx in input directory
        test_doc = Path("input/test_document.docx")
        if test_doc.exists():
            docx_file = str(test_doc)
        else:
            print("Usage: python try_word_to_md.py <docx_file> [output_md_file]")
            return 1
    else:
        docx_file = sys.argv[1]

    output_file = sys.argv[2] if len(sys.argv) > 2 else None

    if not os.path.exists(docx_file):
        print(f"Error: File not found - {docx_file}")
        return 1

    # Create output directory if it doesn't exist
    output_dir = Path("output")
    output_dir.mkdir(exist_ok=True)

    # Default output file if not provided
    if output_file is None:
        output_file = output_dir / Path(docx_file).with_suffix(".md").name

    converted_file = convert_docx_to_markdown(docx_file, output_file)
    print(f"\nSuccess! Converted {docx_file} to {converted_file}")
    print("Now you can use this file with Project Requirements.py")

    return 0


if __name__ == "__main__":
    sys.exit(main())
