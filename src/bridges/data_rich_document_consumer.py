#!/usr/bin/env python3
"""Consume a data-rich document projection and render a governed DOCX.

The upstream JSON stays authoritative. This consumer validates the outward
manifest, renders Markdown through Pandoc + reference.docx, then emits a return
receipt with exact hashes and structural heading/numbering proof.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Tuple
from xml.etree import ElementTree as ET

from docx import Document


MANIFEST_SCHEMA = "gbogeb.docx-rtm-outward-document-manifest/1.0.0"
CONSUMER_REPO = "GBOGEB/DOCX_RTM_Automation"
STYLE_SCHEMA = "docx_rtm.visual_style/1.0.0"
DEFAULT_STYLE_PATH = Path(__file__).resolve().parents[2] / "federation" / "DATA_RICH_DOCUMENT" / "visual_style.json"
NUMBERED_TITLE_RE = re.compile(r"^\s*\d{1,3}(?:\.\d+)*[.)]?\s+")
REQUIREMENT_TITLE_RE = re.compile(r"^\s*REQ-\d+\s+[—-]\s+")
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


class ConsumerError(ValueError):
    pass


def w(tag: str) -> str:
    return f"{{{W}}}{tag}"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_manifest(path: Path) -> Dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def load_visual_style(path: Path | None = None) -> Dict[str, Any]:
    style_path = path or DEFAULT_STYLE_PATH
    try:
        style = json.loads(style_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ConsumerError(f"cannot load visual style {style_path}: {exc}") from exc

    if style.get("schema") != STYLE_SCHEMA:
        raise ConsumerError(
            f"unsupported visual style schema: {style.get('schema')!r}"
        )
    if not style.get("style_id"):
        raise ConsumerError("visual style style_id is required")
    governance = style.get("governance", {})
    if governance.get("semantic_content_change_allowed") is not False:
        raise ConsumerError("visual style must forbid semantic content changes")
    if governance.get("accepted_baseline_mutation_allowed") is not False:
        raise ConsumerError("visual style must forbid accepted baseline mutation")
    return style


def validate_manifest(
    manifest: Dict[str, Any],
    markdown_path: Path,
    expected_source_ref: str,
) -> None:
    if manifest.get("schema") != MANIFEST_SCHEMA:
        raise ConsumerError(
            f"unsupported manifest schema: {manifest.get('schema')!r}"
        )

    authority = manifest.get("authority", {})
    forbidden_true = {
        "rendered_output_authoritative",
        "authority_transfer",
        "engineering_credit",
        "compliance_credit",
        "semantic_changes_allowed_in_consumer",
    }
    for key in forbidden_true:
        if authority.get(key) is not False:
            raise ConsumerError(f"authority guard {key!r} must be false")

    consumer = manifest.get("docx_rtm_consumer", {})
    if consumer.get("repo") != CONSUMER_REPO:
        raise ConsumerError(
            f"manifest consumer repo mismatch: {consumer.get('repo')!r}"
        )
    if consumer.get("forbidden_numbering_filter") != "config/filters/extend_headings.lua":
        raise ConsumerError("manifest must explicitly forbid literal heading-number filter")

    source_ref = manifest.get("source", {}).get("git_ref")
    if not source_ref:
        raise ConsumerError("manifest source.git_ref is required")
    if expected_source_ref and source_ref != expected_source_ref:
        raise ConsumerError(
            f"source git ref mismatch: observed {source_ref}, expected {expected_source_ref}"
        )

    observed_md_hash = sha256_file(markdown_path)
    expected_md_hash = manifest.get("projection", {}).get("sha256")
    if observed_md_hash != expected_md_hash:
        raise ConsumerError(
            f"projection markdown hash mismatch: {observed_md_hash} != {expected_md_hash}"
        )

    heading = manifest.get("projection", {}).get("heading_contract", {})
    if heading.get("numbering") != "TEMPLATE_MANAGED":
        raise ConsumerError("projection numbering must be TEMPLATE_MANAGED")
    if heading.get("heading_text_contains_numbers") is not False:
        raise ConsumerError("projection headings must not contain literal numbers")


def _zip_xml(docx_path: Path, member: str) -> ET.Element:
    with zipfile.ZipFile(docx_path) as archive:
        try:
            payload = archive.read(member)
        except KeyError as exc:
            raise ConsumerError(f"{docx_path}: missing {member}") from exc
    return ET.fromstring(payload)


def inspect_numbering_contract(docx_path: Path) -> Dict[str, Any]:
    styles = _zip_xml(docx_path, "word/styles.xml")
    numbering = _zip_xml(docx_path, "word/numbering.xml")

    style_links: Dict[str, Dict[str, str]] = {}
    for level in (1, 2, 3):
        style_id = f"Heading{level}"
        style = styles.find(f".//w:style[@w:styleId='{style_id}']", NS)
        if style is None:
            raise ConsumerError(f"{docx_path}: missing style {style_id}")

        num_id_el = style.find("./w:pPr/w:numPr/w:numId", NS)
        ilvl_el = style.find("./w:pPr/w:numPr/w:ilvl", NS)
        if num_id_el is None or ilvl_el is None:
            raise ConsumerError(
                f"{docx_path}: {style_id} is not linked to multilevel numbering"
            )

        style_links[style_id] = {
            "num_id": num_id_el.get(w("val"), ""),
            "ilvl": ilvl_el.get(w("val"), ""),
        }

    num_ids = {item["num_id"] for item in style_links.values()}
    if len(num_ids) != 1:
        raise ConsumerError(
            f"{docx_path}: Heading1-3 do not share one numbering instance: {style_links}"
        )
    num_id = next(iter(num_ids))

    num = numbering.find(f".//w:num[@w:numId='{num_id}']", NS)
    if num is None:
        raise ConsumerError(f"{docx_path}: numbering instance {num_id} missing")
    abstract_id_el = num.find("./w:abstractNumId", NS)
    if abstract_id_el is None:
        raise ConsumerError(f"{docx_path}: numbering instance {num_id} has no abstractNum")
    abstract_id = abstract_id_el.get(w("val"), "")

    abstract = numbering.find(
        f".//w:abstractNum[@w:abstractNumId='{abstract_id}']",
        NS,
    )
    if abstract is None:
        raise ConsumerError(
            f"{docx_path}: abstract numbering {abstract_id} missing"
        )

    expected_text = {0: "%1", 1: "%1.%2", 2: "%1.%2.%3"}
    level_details: Dict[str, Dict[str, str]] = {}
    for ilvl, expected_style in enumerate(("Heading1", "Heading2", "Heading3")):
        lvl = abstract.find(f"./w:lvl[@w:ilvl='{ilvl}']", NS)
        if lvl is None:
            raise ConsumerError(f"{docx_path}: numbering level {ilvl} missing")
        pstyle = lvl.find("./w:pStyle", NS)
        lvl_text = lvl.find("./w:lvlText", NS)
        if pstyle is None or pstyle.get(w("val")) != expected_style:
            raise ConsumerError(
                f"{docx_path}: level {ilvl} not linked to {expected_style}"
            )
        if lvl_text is None or lvl_text.get(w("val")) != expected_text[ilvl]:
            raise ConsumerError(
                f"{docx_path}: level {ilvl} numbering text mismatch"
            )
        rpr = lvl.find("./w:rPr", NS)
        number_color_el = rpr.find("./w:color", NS) if rpr is not None else None
        number_bold_el = rpr.find("./w:b", NS) if rpr is not None else None
        number_font_el = rpr.find("./w:rFonts", NS) if rpr is not None else None
        level_details[str(ilvl)] = {
            "pstyle": pstyle.get(w("val"), ""),
            "lvl_text": lvl_text.get(w("val"), ""),
            "number_color": number_color_el.get(w("val"), "") if number_color_el is not None else "",
            "number_bold": number_bold_el is not None,
            "number_font": number_font_el.get(w("ascii"), "") if number_font_el is not None else "",
        }

    for level, style_id in enumerate(("Heading1", "Heading2", "Heading3")):
        if style_links[style_id]["ilvl"] != str(level):
            raise ConsumerError(
                f"{docx_path}: {style_id} ilvl mismatch: {style_links[style_id]}"
            )

    return {
        "status": "PASS",
        "num_id": num_id,
        "abstract_num_id": abstract_id,
        "styles": style_links,
        "levels": level_details,
    }


def _paragraph_style_and_text(p: ET.Element) -> Tuple[str, str]:
    pstyle = p.find("./w:pPr/w:pStyle", NS)
    style = pstyle.get(w("val"), "") if pstyle is not None else ""
    text = "".join(node.text or "" for node in p.findall(".//w:t", NS))
    return style, text


def inspect_rendered_headings(
    docx_path: Path,
    manifest: Dict[str, Any],
) -> Dict[str, Any]:
    document = _zip_xml(docx_path, "word/document.xml")
    observed: List[Tuple[str, str]] = []
    for paragraph in document.findall(".//w:body/w:p", NS):
        style, text = _paragraph_style_and_text(paragraph)
        if style in {"Heading1", "Heading2", "Heading3"}:
            observed.append((style, text))

    expected: List[Tuple[str, str]] = []
    for row in manifest.get("section_map", []):
        level = int(row["heading_level"])
        expected.append((f"Heading{level}", row["title"]))

    if observed != expected:
        raise ConsumerError(
            f"rendered heading sequence mismatch: observed={observed}, expected={expected}"
        )

    numbered = [
        {"style": style, "text": text}
        for style, text in observed
        if NUMBERED_TITLE_RE.match(text)
    ]
    if numbered:
        raise ConsumerError(
            f"literal heading numbering detected in rendered DOCX: {numbered}"
        )

    return {
        "heading_style_check": "PASS",
        "duplicate_numbering_check": "PASS",
        "observed": [
            {"style": style, "text": text}
            for style, text in observed
        ],
    }



def _paragraph_style_name(paragraph) -> str:
    style = paragraph.style
    return style.name if style is not None else ""


def _requirement_blocks(doc: Document) -> List[Tuple[int, int]]:
    """Return inclusive paragraph ranges for rendered requirement blocks."""
    paragraphs = doc.paragraphs
    blocks: List[Tuple[int, int]] = []
    index = 0
    while index < len(paragraphs):
        if not REQUIREMENT_TITLE_RE.match(paragraphs[index].text.strip()):
            index += 1
            continue

        start = index
        end = index
        cursor = index + 1
        while cursor < len(paragraphs):
            text = paragraphs[cursor].text.strip()
            style_name = _paragraph_style_name(paragraphs[cursor])
            if style_name in {"Heading 1", "Heading 2", "Heading 3"}:
                break
            if REQUIREMENT_TITLE_RE.match(text):
                break
            if text:
                end = cursor
            cursor += 1

        blocks.append((start, end))
        index = max(cursor, index + 1)
    return blocks


def enforce_requirement_block_pagination(docx_path: Path) -> Dict[str, Any]:
    """Apply Word paragraph pagination hints to requirement blocks."""
    doc = Document(docx_path)
    blocks = _requirement_blocks(doc)
    for start, end in blocks:
        for index in range(start, end + 1):
            fmt = doc.paragraphs[index].paragraph_format
            fmt.keep_together = True
            fmt.keep_with_next = index < end
    doc.save(docx_path)
    return {
        "status": "PASS",
        "requirement_block_count": len(blocks),
        "mode": "KEEP_WITH_NEXT_HINTS",
    }


def inspect_requirement_pagination(docx_path: Path) -> Dict[str, Any]:
    doc = Document(docx_path)
    blocks = _requirement_blocks(doc)
    if not blocks:
        raise ConsumerError(f"{docx_path}: no governed requirement blocks found")
    observations = []
    for start, end in blocks:
        for index in range(start, end + 1):
            fmt = doc.paragraphs[index].paragraph_format
            if fmt.keep_together is not True:
                raise ConsumerError(
                    f"{docx_path}: requirement paragraph {index} is not keep-together"
                )
            if index < end and fmt.keep_with_next is not True:
                raise ConsumerError(
                    f"{docx_path}: requirement paragraph {index} is not keep-with-next"
                )
        match = re.search(r"REQ-\d+", doc.paragraphs[start].text)
        observations.append(
            {
                "requirement_id": match.group(0) if match else "",
                "paragraph_count": end - start + 1,
            }
        )
    return {
        "status": "PASS",
        "requirement_block_count": len(observations),
        "mode": "KEEP_WITH_NEXT_HINTS",
        "observations": observations,
    }

def _replace_runs(paragraph, segments: List[Tuple[str, str | None]]) -> None:
    """Replace plain runs while preserving paragraph identity and exact text."""
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    for text, style_name in segments:
        if not text:
            continue
        run = paragraph.add_run(text)
        if style_name:
            run.style = style_name


def apply_visual_semantics(
    docx_path: Path,
    visual: Dict[str, Any],
) -> Dict[str, Any]:
    """Apply governed semantic visual roles after Pandoc rendering."""
    doc = Document(docx_path)
    requirements = visual["requirements"]
    captions = visual["captions"]
    metadata_prefixes = tuple(requirements.get("metadata_prefixes", []))

    counts = {
        "requirement_titles": 0,
        "metadata_lines": 0,
        "captions": 0,
    }

    for paragraph in doc.paragraphs:
        text = paragraph.text

        req_match = re.match(r"^\s*(REQ-\d+)(.*)$", text)
        if req_match:
            paragraph.style = doc.styles[requirements["title_paragraph_style"]]
            _replace_runs(
                paragraph,
                [
                    (req_match.group(1), requirements["id_character_style"]),
                    (req_match.group(2), None),
                ],
            )
            counts["requirement_titles"] += 1
            continue

        prefix = next(
            (item for item in metadata_prefixes if text.startswith(item)),
            None,
        )
        if prefix:
            paragraph.style = doc.styles[requirements["metadata_paragraph_style"]]
            _replace_runs(
                paragraph,
                [
                    (prefix, requirements["metadata_label_character_style"]),
                    (text[len(prefix):], None),
                ],
            )
            counts["metadata_lines"] += 1
            continue

        caption_match = re.match(r"^\s*((?:Figure|Table)\s+\d+)(.*)$", text)
        if paragraph.style.name == captions["paragraph_style"] or caption_match:
            paragraph.style = doc.styles[captions["paragraph_style"]]
            if caption_match:
                _replace_runs(
                    paragraph,
                    [
                        (caption_match.group(1), captions["number_character_style"]),
                        (caption_match.group(2), None),
                    ],
                )
            counts["captions"] += 1

    doc.save(docx_path)
    return {"status": "PASS", **counts}


def _style_observation(doc: Document, name: str) -> Dict[str, Any]:
    try:
        style = doc.styles[name]
    except KeyError as exc:
        raise ConsumerError(f"{doc.part.partname}: missing visual style {name!r}") from exc
    rgb = style.font.color.rgb
    return {
        "font": style.font.name or "",
        "size_pt": round(style.font.size.pt, 3) if style.font.size is not None else None,
        "color": str(rgb).upper() if rgb is not None else "",
        "bold": style.font.bold,
        "italic": style.font.italic,
    }


def inspect_visual_style_contract(
    docx_path: Path,
    visual: Dict[str, Any],
) -> Dict[str, Any]:
    doc = Document(docx_path)
    sizes = visual["sizes_pt"]
    colors = visual["colors"]
    requirements = visual["requirements"]
    captions = visual["captions"]

    expected = {
        "Normal": {
            "font": visual["fonts"]["body"]["name"],
            "size_pt": float(sizes["body"]),
            "color": colors["body"].upper(),
        },
        "Heading 1": {
            "font": visual["fonts"]["heading"]["name"],
            "size_pt": float(sizes["heading_1"]),
            "color": colors["heading_primary"].upper(),
        },
        "Heading 2": {
            "font": visual["fonts"]["heading"]["name"],
            "size_pt": float(sizes["heading_2"]),
            "color": colors["heading_secondary"].upper(),
        },
        "Heading 3": {
            "font": visual["fonts"]["heading"]["name"],
            "size_pt": float(sizes["heading_3"]),
            "color": colors["heading_tertiary"].upper(),
        },
        captions["paragraph_style"]: {
            "font": visual["fonts"]["body"]["name"],
            "size_pt": float(sizes["caption"]),
            "color": colors["caption"].upper(),
        },
        requirements["title_paragraph_style"]: {
            "font": visual["fonts"]["heading"]["name"],
            "size_pt": float(sizes["requirement_title"]),
            "color": colors["requirement_title"].upper(),
        },
        requirements["metadata_paragraph_style"]: {
            "font": visual["fonts"]["body"]["name"],
            "size_pt": float(sizes["metadata"]),
            "color": colors["metadata"].upper(),
        },
        requirements["id_character_style"]: {
            "font": visual["fonts"]["heading"]["name"],
            "size_pt": float(sizes["requirement_title"]),
            "color": colors["requirement_id"].upper(),
        },
        requirements["metadata_label_character_style"]: {
            "font": visual["fonts"]["body"]["name"],
            "size_pt": float(sizes["metadata"]),
            "color": colors["special_number"].upper(),
        },
        captions["number_character_style"]: {
            "font": visual["fonts"]["body"]["name"],
            "size_pt": float(sizes["caption"]),
            "color": colors[captions["number_color_role"]].upper(),
        },
    }

    observations: Dict[str, Dict[str, Any]] = {}
    failures: List[str] = []
    for name, wanted in expected.items():
        observed = _style_observation(doc, name)
        observations[name] = observed
        if observed["font"] != wanted["font"]:
            failures.append(f"{name}:font={observed['font']!r}!={wanted['font']!r}")
        if observed["size_pt"] is None or abs(observed["size_pt"] - wanted["size_pt"]) > 0.06:
            failures.append(f"{name}:size={observed['size_pt']}!={wanted['size_pt']}")
        if observed["color"] != wanted["color"]:
            failures.append(f"{name}:color={observed['color']}!={wanted['color']}")

    numbering = inspect_numbering_contract(docx_path)
    expected_number_color = colors[visual["numbering"]["color_role"]].upper()
    for level, details in numbering["levels"].items():
        if details.get("number_color", "").upper() != expected_number_color:
            failures.append(
                f"numbering[{level}]:color={details.get('number_color')!r}!={expected_number_color}"
            )
        if visual["numbering"].get("bold") is True and details.get("number_bold") is not True:
            failures.append(f"numbering[{level}]:bold=false")

    if failures:
        raise ConsumerError("visual style contract failed: " + "; ".join(failures))

    return {
        "status": "PASS",
        "style_id": visual["style_id"],
        "styles": observations,
        "numbering_color": expected_number_color,
    }


def render_docx(
    markdown_path: Path,
    reference_docx: Path,
    output_docx: Path,
) -> None:
    if not reference_docx.exists():
        raise ConsumerError(f"reference DOCX missing: {reference_docx}")

    reference_proof = inspect_numbering_contract(reference_docx)
    if reference_proof["status"] != "PASS":
        raise ConsumerError("reference numbering contract failed")

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    cmd = [
        "pandoc",
        str(markdown_path),
        "-o",
        str(output_docx),
        f"--reference-doc={reference_docx}",
    ]
    result = subprocess.run(cmd, capture_output=True, text=True)
    if result.returncode != 0:
        raise ConsumerError(f"pandoc render failed: {result.stderr.strip()}")
    if not output_docx.exists() or output_docx.stat().st_size == 0:
        raise ConsumerError("pandoc completed without a non-empty DOCX")

    enforce_requirement_block_pagination(output_docx)


def build_return_receipt(
    manifest: Dict[str, Any],
    markdown_path: Path,
    reference_docx: Path,
    output_docx: Path,
) -> Dict[str, Any]:
    numbering = inspect_numbering_contract(output_docx)
    headings = inspect_rendered_headings(output_docx, manifest)
    pagination = inspect_requirement_pagination(output_docx)

    return {
        "schema": "docx_rtm.data_rich_document_render_receipt/1.0.0",
        "source_json_sha256": manifest["source"]["sha256"],
        "projection_markdown_sha256": sha256_file(markdown_path),
        "source_git_ref": manifest["source"]["git_ref"],
        "reference_doc_sha256": sha256_file(reference_docx),
        "rendered_docx_sha256": sha256_file(output_docx),
        "render_status": "PASS",
        "heading_style_check": headings["heading_style_check"],
        "template_numbering_check": numbering["status"],
        "duplicate_numbering_check": headings["duplicate_numbering_check"],
        "requirement_pagination_check": pagination["status"],
        "requirement_pagination_observations": pagination,
        "heading_observations": headings["observed"],
        "numbering_observations": numbering,
        "authority": {
            "authority_transfer": False,
            "rendered_output_authoritative": False,
            "engineering_credit": False,
            "compliance_credit": False,
            "semantic_writeback": "JSON_CHANGE_SET_ONLY",
        },
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--markdown", type=Path, required=True)
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--reference-doc", type=Path, required=True)
    parser.add_argument("--output-docx", type=Path, required=True)
    parser.add_argument("--receipt", type=Path, required=True)
    parser.add_argument("--expected-source-ref", default="")
    args = parser.parse_args()

    try:
        manifest = load_manifest(args.manifest)
        validate_manifest(
            manifest,
            args.markdown,
            expected_source_ref=args.expected_source_ref,
        )
        render_docx(args.markdown, args.reference_doc, args.output_docx)
        receipt = build_return_receipt(
            manifest,
            args.markdown,
            args.reference_doc,
            args.output_docx,
        )
    except (OSError, json.JSONDecodeError, ConsumerError) as exc:
        print(f"FAIL: {exc}")
        return 1

    args.receipt.parent.mkdir(parents=True, exist_ok=True)
    args.receipt.write_text(
        json.dumps(receipt, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    print(f"PASS: rendered DOCX -> {args.output_docx}")
    print(f"PASS: render receipt -> {args.receipt}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
