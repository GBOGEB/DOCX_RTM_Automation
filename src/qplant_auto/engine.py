#!/usr/bin/env python3
"""QPLANT-AI-AUTO structured document iteration engine.

The engine keeps RFO, ADR and OCD as one versioned structured state. New input is
parsed, reconciled against the complete prior state through OpenAI Structured
Outputs, validated locally, diffed, and exported as human- and machine-facing
artifacts.

Authority rules:
- user-mandated VERBATIM / OVERRIDE content is not rewritten;
- immutable clauses are never changed by the model;
- missing fields remain present and may be empty;
- inferred fields are marked INFERRED and require review;
- every iteration emits a complete project state, not a patch-only fragment.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import os
import re
import shutil
from dataclasses import dataclass
from enum import Enum
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

try:
    import pandas as pd
except ImportError:  # pragma: no cover - diagnosed by --check
    pd = None  # type: ignore[assignment]

try:
    from docx import Document
except ImportError:  # pragma: no cover - diagnosed by --check
    Document = None  # type: ignore[assignment]


DEFAULT_MODEL = os.environ.get("OPENAI_MODEL", "gpt-5.6")
PILLARS = ("RFO", "ADR", "OCD")
OUTLINE_RE = re.compile(r"^(?P<num>\d+(?:\.\d+){0,2})\s+(?P<title>.+)$")

REQUIREMENT_FIELDS = (
    "id",
    "title",
    "requirement",
    "measurability",
    "reference_standards",
    "verification",
    "validation",
    "rationale",
    "topic_epic",
    "related_requirements",
    "parent_stakeholder_ids",
    "requirement_type",
    "requirement_category",
    "source_refs",
    "status",
    "confidence",
    "notes",
    "field_provenance",
    "immutable",
    "needs_review",
)

SYSTEM_INSTRUCTIONS = """
You are the governed integration engine for QPLANT procurement documentation.
Maintain one complete structured state spanning RFO, ADR and OCD.

Your job is to structure, combine, refine and expand content from the prior
baseline and the new user-owned input while preserving traceability.

Rules:
1. Return the COMPLETE project state on every iteration. Never return a patch.
2. Preserve every schema field. Unknown information is an empty string, empty
   array, or explicit unresolved question; never omit a field.
3. Preserve outline hierarchy through level 3 only (e.g. 3, 3.2, 3.2.1).
4. RFO, ADR and OCD must remain cross-aligned. A change in one pillar must be
   evaluated for required impacts in the other pillars and in the RTM.
5. Requirement IDs and stakeholder IDs are stable identities. Do not renumber
   existing IDs merely because items move.
6. A new requirement may arrive with only a title. You may draft missing fields
   from explicit context or closely related baseline requirements, but every
   inferred field must be marked INFERRED in field_provenance and needs_review
   must be true. Do not invent standards, numeric limits, acceptance criteria or
   engineering facts that are not supported.
7. USER_MANDATED / VERBATIM / OVERRIDE input has model-owner priority over
   non-immutable prior text. VERBATIM text must be retained exactly.
8. Immutable clauses must remain byte-for-byte unchanged in their body text.
9. Preserve superseded material through change history/source references rather
   than silently deleting it.
10. Distinguish source-bound facts from inferred drafting. Empty is preferable
    to fabricated certainty.
11. Keep explanatory prose and machine fields consistent.
12. The output must satisfy the supplied JSON schema exactly.
""".strip()


class InputIntent(str, Enum):
    NEW = "NEW"
    UPDATE = "UPDATE"
    REFINE = "REFINE"
    VERBATIM = "VERBATIM"
    OVERRIDE = "OVERRIDE"
    BRANCH = "BRANCH"
    MASTER = "MASTER"


@dataclass(frozen=True)
class InputEnvelope:
    path: Path
    intent: InputIntent
    source: str
    target_pillars: tuple[str, ...]
    dependencies: tuple[str, ...]
    instructions: str
    branch_name: str = ""


def _string_array() -> dict[str, Any]:
    return {"type": "array", "items": {"type": "string"}}


def _field_provenance_schema() -> dict[str, Any]:
    return {
        "type": "array",
        "items": {
            "type": "object",
            "additionalProperties": False,
            "required": ["field", "source_state", "basis"],
            "properties": {
                "field": {"type": "string"},
                "source_state": {
                    "type": "string",
                    "enum": ["SOURCE_BOUND", "INFERRED", "USER_MANDATED", "EMPTY"],
                },
                "basis": {"type": "string"},
            },
        },
    }


REQUIREMENT_SCHEMA: dict[str, Any] = {
    "type": "object",
    "additionalProperties": False,
    "required": list(REQUIREMENT_FIELDS),
    "properties": {
        "id": {"type": "string"},
        "title": {"type": "string"},
        "requirement": {"type": "string"},
        "measurability": {"type": "string"},
        "reference_standards": _string_array(),
        "verification": {"type": "string"},
        "validation": {"type": "string"},
        "rationale": {"type": "string"},
        "topic_epic": {"type": "string"},
        "related_requirements": _string_array(),
        "parent_stakeholder_ids": _string_array(),
        "requirement_type": {"type": "string"},
        "requirement_category": {"type": "string"},
        "source_refs": _string_array(),
        "status": {"type": "string"},
        "confidence": {"type": "string"},
        "notes": {"type": "string"},
        "field_provenance": _field_provenance_schema(),
        "immutable": {"type": "boolean"},
        "needs_review": {"type": "boolean"},
    },
}

SECTION_SCHEMA: dict[str, Any] = {
    "type": "object",
    "additionalProperties": False,
    "required": [
        "pillar",
        "outline_number",
        "level",
        "title",
        "body_text",
        "parent_outline",
        "requirement_ids",
        "source_refs",
        "immutable",
        "status",
    ],
    "properties": {
        "pillar": {"type": "string", "enum": list(PILLARS)},
        "outline_number": {"type": "string"},
        "level": {"type": "integer", "minimum": 1, "maximum": 3},
        "title": {"type": "string"},
        "body_text": {"type": "string"},
        "parent_outline": {"type": "string"},
        "requirement_ids": _string_array(),
        "source_refs": _string_array(),
        "immutable": {"type": "boolean"},
        "status": {"type": "string"},
    },
}

PROJECT_STATE_SCHEMA: dict[str, Any] = {
    "type": "object",
    "additionalProperties": False,
    "required": [
        "schema_version",
        "project_name",
        "model_name",
        "baseline_version",
        "iteration_id",
        "documents",
        "requirements",
        "relationships",
        "unresolved_questions",
        "change_log",
    ],
    "properties": {
        "schema_version": {"type": "string"},
        "project_name": {"type": "string"},
        "model_name": {"type": "string"},
        "baseline_version": {"type": "string"},
        "iteration_id": {"type": "string"},
        "documents": {
            "type": "object",
            "additionalProperties": False,
            "required": list(PILLARS),
            "properties": {
                pillar: {
                    "type": "object",
                    "additionalProperties": False,
                    "required": ["title", "version", "sections"],
                    "properties": {
                        "title": {"type": "string"},
                        "version": {"type": "string"},
                        "sections": {"type": "array", "items": SECTION_SCHEMA},
                    },
                }
                for pillar in PILLARS
            },
        },
        "requirements": {"type": "array", "items": REQUIREMENT_SCHEMA},
        "relationships": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": [
                    "from_id",
                    "to_id",
                    "relation_type",
                    "rationale",
                    "source_state",
                ],
                "properties": {
                    "from_id": {"type": "string"},
                    "to_id": {"type": "string"},
                    "relation_type": {"type": "string"},
                    "rationale": {"type": "string"},
                    "source_state": {
                        "type": "string",
                        "enum": ["SOURCE_BOUND", "INFERRED", "USER_MANDATED", "EMPTY"],
                    },
                },
            },
        },
        "unresolved_questions": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["id", "question", "affected_ids", "status"],
                "properties": {
                    "id": {"type": "string"},
                    "question": {"type": "string"},
                    "affected_ids": _string_array(),
                    "status": {"type": "string"},
                },
            },
        },
        "change_log": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": [
                    "iteration_id",
                    "timestamp",
                    "input_sha256",
                    "intent",
                    "summary",
                    "added_ids",
                    "modified_ids",
                    "superseded_ids",
                    "unchanged_count",
                ],
                "properties": {
                    "iteration_id": {"type": "string"},
                    "timestamp": {"type": "string"},
                    "input_sha256": {"type": "string"},
                    "intent": {"type": "string"},
                    "summary": {"type": "string"},
                    "added_ids": _string_array(),
                    "modified_ids": _string_array(),
                    "superseded_ids": _string_array(),
                    "unchanged_count": {"type": "integer", "minimum": 0},
                },
            },
        },
    },
}


def empty_project_state(project_name: str = "QPLANT Procurement") -> dict[str, Any]:
    return {
        "schema_version": "qplant.project_state/v1",
        "project_name": project_name,
        "model_name": "QPLANT-AI-AUTO",
        "baseline_version": "v0.0",
        "iteration_id": "",
        "documents": {
            pillar: {"title": pillar, "version": "v0.0", "sections": []}
            for pillar in PILLARS
        },
        "requirements": [],
        "relationships": [],
        "unresolved_questions": [],
        "change_log": [],
    }


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_project_state(path: Path | None) -> dict[str, Any]:
    if path is None or not path.exists():
        return empty_project_state()
    data = json.loads(path.read_text(encoding="utf-8"))
    validate_project_state(data)
    return data


def _read_docx(path: Path) -> str:
    if Document is None:
        raise RuntimeError("python-docx is not installed")
    doc = Document(path)
    blocks: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            style = paragraph.style.name if paragraph.style else ""
            blocks.append(f"[{style}] {text}" if style else text)
    for table_index, table in enumerate(doc.tables, start=1):
        blocks.append(f"[TABLE {table_index}]")
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(blocks)


def read_input(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _read_docx(path)
    if suffix in {".txt", ".md", ".rst"}:
        return path.read_text(encoding="utf-8")
    if suffix == ".json":
        return json.dumps(json.loads(path.read_text(encoding="utf-8")), indent=2)
    if suffix == ".csv":
        if pd is None:
            raise RuntimeError("pandas is not installed")
        return pd.read_csv(path).to_json(orient="records", indent=2)
    if suffix in {".xlsx", ".xlsm"}:
        if pd is None:
            raise RuntimeError("pandas/openpyxl is not installed")
        book = pd.ExcelFile(path)
        payload = {
            sheet: pd.read_excel(path, sheet_name=sheet).fillna("").to_dict(orient="records")
            for sheet in book.sheet_names
        }
        return json.dumps(payload, indent=2, default=str)
    raise ValueError(f"Unsupported input type: {path.suffix}")


def validate_project_state(state: Mapping[str, Any]) -> None:
    missing = [key for key in PROJECT_STATE_SCHEMA["required"] if key not in state]
    if missing:
        raise ValueError(f"Project state missing top-level fields: {missing}")

    seen_req: set[str] = set()
    for requirement in state.get("requirements", []):
        missing_req = [key for key in REQUIREMENT_FIELDS if key not in requirement]
        if missing_req:
            raise ValueError(
                f"Requirement {requirement.get('id', '<unknown>')} missing fields: {missing_req}"
            )
        req_id = requirement["id"]
        if not req_id:
            raise ValueError("Requirement id cannot be empty")
        if req_id in seen_req:
            raise ValueError(f"Duplicate requirement id: {req_id}")
        seen_req.add(req_id)

    for pillar in PILLARS:
        document = state["documents"].get(pillar)
        if not isinstance(document, Mapping):
            raise ValueError(f"Missing document pillar: {pillar}")
        seen_outline: set[str] = set()
        for section in document.get("sections", []):
            number = section.get("outline_number", "")
            if not OUTLINE_RE.match(f"{number} x"):
                raise ValueError(f"Invalid outline number {number!r} in {pillar}")
            level = len(number.split("."))
            if level != section.get("level") or level > 3:
                raise ValueError(
                    f"Outline/level mismatch in {pillar}: {number} / {section.get('level')}"
                )
            if number in seen_outline:
                raise ValueError(f"Duplicate outline number {pillar} {number}")
            seen_outline.add(number)


class QPlantOpenAI:
    """OpenAI Responses API adapter using strict JSON-schema Structured Outputs."""

    def __init__(
        self,
        api_key: str | None = None,
        api_key_file: Path | None = None,
        model: str = DEFAULT_MODEL,
    ) -> None:
        self.model = model
        self.api_key = api_key or os.environ.get("OPENAI_API_KEY")
        if not self.api_key and api_key_file:
            self.api_key = api_key_file.expanduser().read_text(encoding="utf-8").strip()
        self._client: Any = None

    @property
    def client(self) -> Any:
        if self._client is None:
            if not self.api_key:
                raise RuntimeError(
                    "OPENAI_API_KEY is not set. Prefer an environment variable; "
                    "alternatively pass --api-key-file pointing outside the repository."
                )
            try:
                from openai import OpenAI
            except ImportError as exc:  # pragma: no cover
                raise RuntimeError("Install the OpenAI SDK: python -m pip install openai") from exc
            self._client = OpenAI(api_key=self.api_key)
        return self._client

    def probe(self) -> str:
        response = self.client.responses.create(
            model=self.model,
            input="Reply with exactly QPLANT_API_OK.",
            max_output_tokens=32,
        )
        return response.output_text.strip()

    def reconcile(
        self,
        baseline: Mapping[str, Any],
        envelope: InputEnvelope,
        input_text: str,
        *,
        iteration_id: str,
        timestamp: str,
        input_sha256: str,
    ) -> dict[str, Any]:
        prompt_payload = {
            "iteration": {
                "iteration_id": iteration_id,
                "timestamp": timestamp,
                "intent": envelope.intent.value,
                "source": envelope.source,
                "input_path": envelope.path.name,
                "input_sha256": input_sha256,
                "target_pillars": list(envelope.target_pillars),
                "dependencies": list(envelope.dependencies),
                "instructions": envelope.instructions,
                "branch_name": envelope.branch_name,
            },
            "baseline_project_state": baseline,
            "new_input": input_text,
        }
        response = self.client.responses.create(
            model=self.model,
            instructions=SYSTEM_INSTRUCTIONS,
            input=json.dumps(prompt_payload, ensure_ascii=False),
            text={
                "format": {
                    "type": "json_schema",
                    "name": "qplant_project_state",
                    "description": (
                        "Complete governed RFO/ADR/OCD project state and RTM after one iteration"
                    ),
                    "strict": True,
                    "schema": PROJECT_STATE_SCHEMA,
                }
            },
            max_output_tokens=30000,
        )
        if not response.output_text:
            raise RuntimeError("OpenAI response contained no output_text")
        state = json.loads(response.output_text)
        validate_project_state(state)
        return state


def _requirement_map(state: Mapping[str, Any]) -> dict[str, Mapping[str, Any]]:
    return {item["id"]: item for item in state.get("requirements", [])}


def _section_map(state: Mapping[str, Any]) -> dict[str, Mapping[str, Any]]:
    result: dict[str, Mapping[str, Any]] = {}
    for pillar in PILLARS:
        for section in state["documents"][pillar]["sections"]:
            result[f"{pillar}:{section['outline_number']}"] = section
    return result


def compute_diff(before: Mapping[str, Any], after: Mapping[str, Any]) -> dict[str, Any]:
    before_items: dict[str, Mapping[str, Any]] = {}
    before_items.update(_requirement_map(before))
    before_items.update(_section_map(before))
    after_items: dict[str, Mapping[str, Any]] = {}
    after_items.update(_requirement_map(after))
    after_items.update(_section_map(after))

    before_ids = set(before_items)
    after_ids = set(after_items)
    added = sorted(after_ids - before_ids)
    superseded = sorted(before_ids - after_ids)
    modified = sorted(
        key for key in before_ids & after_ids if before_items[key] != after_items[key]
    )
    unchanged = len((before_ids & after_ids) - set(modified))
    return {
        "added_ids": added,
        "modified_ids": modified,
        "superseded_ids": superseded,
        "unchanged_count": unchanged,
    }


def enforce_immutable(before: Mapping[str, Any], after: Mapping[str, Any]) -> None:
    before_sections = _section_map(before)
    after_sections = _section_map(after)
    for key, old in before_sections.items():
        if not old.get("immutable"):
            continue
        new = after_sections.get(key)
        if new is None:
            raise ValueError(f"Immutable section removed: {key}")
        if old.get("body_text") != new.get("body_text"):
            raise ValueError(f"Immutable section changed: {key}")

    before_requirements = _requirement_map(before)
    after_requirements = _requirement_map(after)
    for key, old in before_requirements.items():
        if not old.get("immutable"):
            continue
        new = after_requirements.get(key)
        if new is None:
            raise ValueError(f"Immutable requirement removed: {key}")
        if old != new:
            raise ValueError(f"Immutable requirement changed: {key}")


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False, sort_keys=False) + "\n",
        encoding="utf-8",
    )


def _write_rtm_csv(path: Path, requirements: Sequence[Mapping[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    scalar_fields = [
        "id",
        "title",
        "requirement",
        "measurability",
        "verification",
        "validation",
        "rationale",
        "topic_epic",
        "requirement_type",
        "requirement_category",
        "status",
        "confidence",
        "immutable",
        "needs_review",
    ]
    list_fields = [
        "reference_standards",
        "related_requirements",
        "parent_stakeholder_ids",
        "source_refs",
    ]
    fieldnames = scalar_fields + list_fields
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for requirement in requirements:
            row = {key: requirement.get(key, "") for key in scalar_fields}
            row.update(
                {key: " | ".join(requirement.get(key, [])) for key in list_fields}
            )
            writer.writerow(row)


def _write_rtm_xlsx(path: Path, requirements: Sequence[Mapping[str, Any]]) -> None:
    if pd is None:
        return
    rows: list[dict[str, Any]] = []
    for requirement in requirements:
        row = dict(requirement)
        for key, value in list(row.items()):
            if isinstance(value, list):
                if key == "field_provenance":
                    row[key] = json.dumps(value, ensure_ascii=False)
                else:
                    row[key] = " | ".join(str(item) for item in value)
        rows.append(row)
    pd.DataFrame(rows).to_excel(path, index=False)


def _write_markdown(path: Path, state: Mapping[str, Any]) -> None:
    lines = [
        f"# {state['project_name']}",
        "",
        f"Model: {state['model_name']}",
        f"Baseline: {state['baseline_version']}",
        f"Iteration: {state['iteration_id']}",
        "",
    ]
    for pillar in PILLARS:
        doc = state["documents"][pillar]
        lines.extend([f"# {pillar}: {doc['title']}", ""])
        for section in sorted(
            doc["sections"],
            key=lambda item: [int(part) for part in item["outline_number"].split(".")],
        ):
            depth = min(6, section["level"] + 1)
            lines.extend(
                [
                    f"{'#' * depth} {section['outline_number']} {section['title']}",
                    "",
                    section["body_text"],
                    "",
                ]
            )
    lines.extend(["# Requirements Traceability Matrix", ""])
    for req in state["requirements"]:
        lines.extend(
            [
                f"## {req['id']} — {req['title']}",
                "",
                req["requirement"],
                "",
                f"- Type: {req['requirement_type']}",
                f"- Category: {req['requirement_category']}",
                f"- Parent stakeholder IDs: {', '.join(req['parent_stakeholder_ids'])}",
                f"- Related requirements: {', '.join(req['related_requirements'])}",
                f"- Measurability: {req['measurability']}",
                f"- Verification: {req['verification']}",
                f"- Validation: {req['validation']}",
                f"- Rationale: {req['rationale']}",
                f"- Needs review: {req['needs_review']}",
                "",
            ]
        )
    path.write_text("\n".join(lines), encoding="utf-8")


def _write_docx(path: Path, state: Mapping[str, Any]) -> None:
    if Document is None:
        return
    doc = Document()
    doc.add_heading(state["project_name"], level=0)
    doc.add_paragraph(
        f"Model {state['model_name']} | Baseline {state['baseline_version']} | "
        f"Iteration {state['iteration_id']}"
    )
    for pillar in PILLARS:
        pillar_doc = state["documents"][pillar]
        doc.add_heading(f"{pillar}: {pillar_doc['title']}", level=1)
        for section in sorted(
            pillar_doc["sections"],
            key=lambda item: [int(part) for part in item["outline_number"].split(".")],
        ):
            level = min(3, section["level"] + 1)
            doc.add_heading(
                f"{section['outline_number']} {section['title']}",
                level=level,
            )
            doc.add_paragraph(section["body_text"])
    doc.add_heading("Requirements Traceability Matrix", level=1)
    for req in state["requirements"]:
        doc.add_heading(f"{req['id']} — {req['title']}", level=2)
        doc.add_paragraph(req["requirement"])
        for label, key in (
            ("Measurability", "measurability"),
            ("Verification", "verification"),
            ("Validation", "validation"),
            ("Rationale", "rationale"),
            ("Type", "requirement_type"),
            ("Category", "requirement_category"),
        ):
            doc.add_paragraph(f"{label}: {req[key]}")
    doc.save(path)


def append_input_ledger(
    ledger_path: Path,
    envelope: InputEnvelope,
    *,
    sha256: str,
    iteration_id: str,
    timestamp: str,
) -> None:
    ledger_path.parent.mkdir(parents=True, exist_ok=True)
    entry = {
        "iteration_id": iteration_id,
        "timestamp": timestamp,
        "filename": envelope.path.name,
        "sha256": sha256,
        "intent": envelope.intent.value,
        "source": envelope.source,
        "target_pillars": list(envelope.target_pillars),
        "dependencies": list(envelope.dependencies),
        "instructions": envelope.instructions,
        "branch_name": envelope.branch_name,
    }
    with ledger_path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(entry, ensure_ascii=False) + "\n")


class QPlantIterationEngine:
    def __init__(
        self,
        root: Path,
        ai: QPlantOpenAI,
    ) -> None:
        self.root = root
        self.ai = ai

    def run(
        self,
        baseline_path: Path | None,
        envelope: InputEnvelope,
    ) -> Path:
        baseline = load_project_state(baseline_path)
        input_text = read_input(envelope.path)
        input_hash = sha256_file(envelope.path)
        now = dt.datetime.now(dt.timezone.utc)
        timestamp = now.isoformat()
        iteration_id = now.strftime("%Y%m%dT%H%M%SZ")

        state = self.ai.reconcile(
            baseline,
            envelope,
            input_text,
            iteration_id=iteration_id,
            timestamp=timestamp,
            input_sha256=input_hash,
        )
        enforce_immutable(baseline, state)
        diff = compute_diff(baseline, state)

        if not state["change_log"] or state["change_log"][-1]["iteration_id"] != iteration_id:
            state["change_log"].append(
                {
                    "iteration_id": iteration_id,
                    "timestamp": timestamp,
                    "input_sha256": input_hash,
                    "intent": envelope.intent.value,
                    "summary": "Local diff receipt appended after structured reconciliation.",
                    **diff,
                }
            )
        else:
            state["change_log"][-1].update(diff)
            state["change_log"][-1]["input_sha256"] = input_hash
            state["change_log"][-1]["intent"] = envelope.intent.value
            state["change_log"][-1]["timestamp"] = timestamp

        state["iteration_id"] = iteration_id
        validate_project_state(state)

        version_dir = self.root / "output" / "qplant_ai_auto" / iteration_id
        version_dir.mkdir(parents=True, exist_ok=False)

        _write_json(version_dir / "Project_Data.json", state)
        _write_json(version_dir / "RTM.json", {"requirements": state["requirements"]})
        _write_json(version_dir / "Change_Log.json", {"change_log": state["change_log"]})
        _write_json(
            version_dir / "Input_Receipt.json",
            {
                "iteration_id": iteration_id,
                "source_file": envelope.path.name,
                "sha256": input_hash,
                "intent": envelope.intent.value,
                "target_pillars": list(envelope.target_pillars),
                "dependencies": list(envelope.dependencies),
                "branch_name": envelope.branch_name,
            },
        )
        _write_rtm_csv(version_dir / "RTM.csv", state["requirements"])
        _write_rtm_xlsx(version_dir / "RTM.xlsx", state["requirements"])
        _write_markdown(version_dir / "Consolidated_Project_Document.md", state)
        _write_docx(version_dir / "Consolidated_Project_Document.docx", state)

        latest_dir = self.root / "output" / "qplant_ai_auto" / "LATEST"
        if latest_dir.exists():
            shutil.rmtree(latest_dir)
        shutil.copytree(version_dir, latest_dir)

        append_input_ledger(
            self.root / "state" / "qplant_input_ledger.jsonl",
            envelope,
            sha256=input_hash,
            iteration_id=iteration_id,
            timestamp=timestamp,
        )
        return version_dir


def package_checks() -> dict[str, str]:
    checks: dict[str, str] = {}
    for import_name in ("openai", "pandas", "docx", "openpyxl"):
        try:
            module = __import__(import_name)
            checks[import_name] = getattr(module, "__version__", "installed")
        except Exception as exc:  # pragma: no cover
            checks[import_name] = f"MISSING: {exc}"
    return checks


def _parse_pillars(raw: str) -> tuple[str, ...]:
    values = tuple(item.strip().upper() for item in raw.split(",") if item.strip())
    invalid = sorted(set(values) - set(PILLARS))
    if invalid:
        raise argparse.ArgumentTypeError(f"Unknown pillar(s): {', '.join(invalid)}")
    return values or PILLARS


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="QPLANT-AI-AUTO governed iteration")
    parser.add_argument("--root", type=Path, default=Path.cwd())
    parser.add_argument("--check", action="store_true", help="Check required Python packages")
    parser.add_argument("--probe-api", action="store_true", help="Probe OpenAI API connectivity")
    parser.add_argument("--input", type=Path)
    parser.add_argument("--baseline", type=Path)
    parser.add_argument("--intent", choices=[item.value for item in InputIntent], default="NEW")
    parser.add_argument("--source", default="user_input")
    parser.add_argument("--target-pillars", type=_parse_pillars, default=PILLARS)
    parser.add_argument("--dependencies", default="")
    parser.add_argument("--instructions", default="Integrate and reconcile with the complete project state.")
    parser.add_argument("--branch-name", default="")
    parser.add_argument("--api-key-file", type=Path)
    parser.add_argument("--model", default=DEFAULT_MODEL)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_arg_parser().parse_args(argv)
    if args.check:
        print(json.dumps(package_checks(), indent=2))
        return 0

    ai = QPlantOpenAI(api_key_file=args.api_key_file, model=args.model)
    if args.probe_api:
        result = ai.probe()
        print(result)
        return 0 if result == "QPLANT_API_OK" else 2

    if not args.input:
        raise SystemExit("--input is required unless --check or --probe-api is used")

    target_pillars = args.target_pillars
    if isinstance(target_pillars, str):
        target_pillars = _parse_pillars(target_pillars)
    dependencies = tuple(
        item.strip() for item in args.dependencies.split(",") if item.strip()
    )
    envelope = InputEnvelope(
        path=args.input,
        intent=InputIntent(args.intent),
        source=args.source,
        target_pillars=tuple(target_pillars),
        dependencies=dependencies,
        instructions=args.instructions,
        branch_name=args.branch_name,
    )
    engine = QPlantIterationEngine(args.root, ai)
    output = engine.run(args.baseline, envelope)
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
